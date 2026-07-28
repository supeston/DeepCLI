import asyncio
import concurrent.futures
import ctypes
import difflib
from datetime import datetime, timezone
from importlib.metadata import PackageNotFoundError, version as package_version
import json
import os
import py_compile
import re
import shutil
import signal
import subprocess
import sys
import textwrap
import unicodedata
from pathlib import Path

import pyperclip
from prompt_toolkit import PromptSession
from prompt_toolkit.document import Document
from prompt_toolkit.completion import Completer, Completion
from prompt_toolkit.shortcuts import radiolist_dialog
from prompt_toolkit.styles import Style as PTStyle

from rich.align import Align
from rich.console import Console, Group
from rich.live import Live
from rich.markdown import Markdown
from rich.panel import Panel
from rich.table import Table
from rich.text import Text
from rich import box

from deep_api import DeepAPI
from dadata_osint import DadataOSINTTool
from funstat_osint import FunstatOSINTTool

                                                                                
                    
                                                                                

                                                                          
MAX_READ_LINES = 1200
                                                                          
MAX_TOOL_FEEDBACK_CHARS = 60000
                                                                                         
TODO_NUDGE_AFTER_TOOLS = 4
                                                                                         
MAX_EMPTY_RESPONSE_RETRIES = 2
# The conversation already contains the original request. Repeat it only
# occasionally during unusually long autonomous runs to prevent goal drift.
TASK_GOAL_REMINDER_INTERVAL = 8
VOICE_HOTKEY_VK = 0xA3  # Right Ctrl
VOICE_HOTKEY_LABEL = "правый Ctrl"
VOICE_HOLD_SECONDS = 0.5
VOICE_SAMPLE_RATE = 16000
VOICE_PARTIAL_INTERVAL = 1.25
VOICE_MODEL_NAME = os.environ.get("DEEPX_VOICE_MODEL", "base")

# Agent execution discipline. These rules are injected by the runner and do
# not modify the editable base prompts.
AGENT_EXECUTION_PROTOCOL = """
АГЕНТНЫЙ ЦИКЛ (ОБЯЗАТЕЛЬНО):
1. Перед каждым действием определи: что сейчас неизвестно, какое действие это
   проверит и какой наблюдаемый результат будет доказательством успеха.
2. После каждого результата инструмента сравни ожидание с фактическим выводом.
   Не считай предположение подтверждённым без наблюдаемого результата.
3. Код или файл, который только записан, ещё не готов. Перед завершением выполни
   релевантную проверку результата: тест, запуск, чтение, инспекцию или иную
   проверку, подходящую исходной задаче.
4. Две одинаковые неудачи подряд означают, что нужно изменить способ решения,
   а не повторять то же действие.
5. Финальный ответ разрешён только когда исходная цель выполнена и это
   подтверждается результатами инструментов. Если проверка невозможна, честно
   назови непроверенную часть вместо заявления об успехе.
Не пересказывай внутренний ход рассуждений. Показывай пользователю действия,
проверяемые результаты и выводы.
""".strip()

MUTATING_TOOLS_REQUIRING_VERIFICATION = {
    "write_file",
    "edit_file",
    "make_excel",
    "make_docx",
    "make_pptx",
    "zip_pack",
    "unzip_pack",
}

VERIFICATION_TOOLS = {
    "read_file",
    "file_info",
    "run_cmd",
    "run_python",
    "task_status",
    "task_log",
    "browser_action",
    "fetch_url",
}

# Calls in the same contiguous read-only batch can safely run concurrently.
# Mutating tools remain ordering barriers so a later read observes earlier writes.
PARALLEL_READ_ONLY_TOOLS = {
    "read_file",
    "web_search",
    "fetch_url",
    "dadata_osint",
    "list_dir",
}

MAX_BACKGROUND_LOG_CHARS = 1_000_000
FILE_PROGRESS_MIN_SECONDS = 0.32
FILE_PROGRESS_FRAME_SECONDS = 0.08
FILE_PROGRESS_FRAMES = ("|", "/", "-", "\\")
FILE_VISUAL_TOOL_NAMES = {
    "write_file",
    "edit_file",
    "make_excel",
    "make_docx",
    "make_pptx",
}

# Keep the response visibly streaming one character at a time.  The delay is
# short enough to stay ahead of the model while still producing a smooth
# typewriter effect instead of printing browser polling batches.
STREAM_CHAR_DELAY = 0.004
STREAM_PUNCTUATION_DELAY = 0.008

                                                                                
                                                                    
                                                                                

APP_DIR = os.path.dirname(os.path.abspath(__file__))

                                            
SELF_PROTECTED_FILES = {
    "deep_cli.py",
    "deep_api.py",
    "dadata_osint.py",
    "funstat_osint.py",
    "run_cli.vbs",
}
                                                                                   
SELF_PROTECTED_PREFIXES = ("system_prompt", ".deepx", ".env", "state.json")
                                         
SELF_PROTECTED_DIRS = {"__pycache__"}


def default_workspace_dir() -> str:
                                                                         
    home = os.path.expanduser("~")
    candidates = [
        os.path.join(home, "Desktop"),
        os.path.join(home, "OneDrive", "Desktop"),                      
        os.path.join(home, "Рабочий стол"),
    ]
    for candidate in candidates:
        if os.path.isdir(candidate):
            return os.path.abspath(candidate)
    return os.path.abspath(home)

                                                                                
                                            
                                                                                

_console_ctrl_handler = None


def _begin_windows_ctrl_c_key_mode():
    if sys.platform != "win32":
        return None
    kernel32 = ctypes.windll.kernel32
    handle = kernel32.GetStdHandle(-10)
    if handle in (-1, 0):
        return None
    mode = ctypes.c_ulong()
    if not kernel32.GetConsoleMode(handle, ctypes.byref(mode)):
        return None
    if not kernel32.SetConsoleMode(handle, mode.value & ~0x0001):
        return None
    return handle, mode.value


def _restore_windows_console_mode(state):
    if state and sys.platform == "win32":
        ctypes.windll.kernel32.SetConsoleMode(state[0], state[1])


async def _wait_for_windows_control_key():
    import msvcrt

    while True:
        if msvcrt.kbhit():
            char = msvcrt.getwch()
            if char == "\x03":
                return "exit"
            if char == "\x18":
                return "stop"
            if char in ("\x00", "\xe0") and msvcrt.kbhit():
                msvcrt.getwch()
        await asyncio.sleep(0.02)

def setup_windows_terminal(cols: int = 140, lines: int = 40, buffer_lines: int = 9999):
                                                                                              
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    if hasattr(sys.stderr, 'reconfigure'):
        sys.stderr.reconfigure(encoding='utf-8', errors='replace')

    if sys.platform == "win32":
        try:
            kernel32 = ctypes.windll.kernel32
            hOut = kernel32.GetStdHandle(-11)                     
            hIn = kernel32.GetStdHandle(-10)                     
            if hOut != -1 and hOut != 0:
                mode = ctypes.c_ulong()
                if kernel32.GetConsoleMode(hOut, ctypes.byref(mode)):
                    mode.value |= 0x0004                                      
                    kernel32.SetConsoleMode(hOut, mode)

                if hIn != -1 and hIn != 0:
                    in_mode = ctypes.c_ulong()
                    if kernel32.GetConsoleMode(hIn, ctypes.byref(in_mode)):
                        in_mode.value |= 0x0010                      
                        kernel32.SetConsoleMode(hIn, in_mode)

                if not os.environ.get("WT_SESSION"):
                    class COORD(ctypes.Structure):
                        _fields_ = [("X", ctypes.c_short), ("Y", ctypes.c_short)]
                    class SMALL_RECT(ctypes.Structure):
                        _fields_ = [("Left", ctypes.c_short), ("Top", ctypes.c_short), ("Right", ctypes.c_short), ("Bottom", ctypes.c_short)]

                    coord = COORD(cols, buffer_lines)
                    kernel32.SetConsoleScreenBufferSize(hOut, coord)
                    rect = SMALL_RECT(0, 0, cols - 1, lines - 1)
                    kernel32.SetConsoleWindowInfo(hOut, True, ctypes.byref(rect))

                                                                        
                                                                                         
            kernel32.GetCurrentProcess.restype = ctypes.c_void_p
            kernel32.TerminateProcess.argtypes = [ctypes.c_void_p, ctypes.c_uint]
            kernel32.TerminateProcess.restype = ctypes.c_bool
            PHANDLER_ROUTINE = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_ulong)
            def on_console_close(ctrl_type):
                if ctrl_type in (2, 5, 6):                                                            
                    process = kernel32.GetCurrentProcess()
                    kernel32.TerminateProcess(process, 0)
                    return True
                return False

            global _console_ctrl_handler
            _console_ctrl_handler = PHANDLER_ROUTINE(on_console_close)
            kernel32.SetConsoleCtrlHandler(_console_ctrl_handler, True)

        except Exception:
            try:
                if not os.environ.get("WT_SESSION"):
                    os.system(f"mode con: cols={cols}")
            except Exception:
                pass

setup_windows_terminal(cols=140, lines=40, buffer_lines=9999)
console = Console()

                                                                                
                            
                                                                                

COMMANDS_DICT = {
    "mode": "Switch model mode (Instant, Expert, Vision) -> Auto-creates new chat",
    "style": "Switch agent style and toolset (Coder, Universal, OSINTer)",
    "think": "Toggle DeepThink reasoning engine (ON / OFF)",
    "search": "Toggle Smart Web Search (ON / OFF) [Instant mode only]",
    "menu": "Open interactive TUI settings menu",
    "ctx": "Compress the current task and continue it in a fresh chat",
    "new": "Start a fresh session context",
    "status": "Display active session parameters",
    "stealth": "Check playwright-stealth version and browser fingerprint signals",
    "copy": "Copy last response to clipboard",
    "clear": "Clear terminal screen",
    "help": "Display command list and shortcuts",
    "exit": "Quit DEEPX CLI"
}

class SlashCommandCompleter(Completer):
                                                                     
    def __init__(self, commands: dict):
        self.commands = commands

    def get_completions(self, document, complete_event):
        text = document.text_before_cursor
        if text.startswith("/"):
            query = text[1:].lower()
            for cmd, desc in self.commands.items():
                if cmd.lower().startswith(query):
                    yield Completion(
                        f"/{cmd}",
                        start_position=-len(text),
                        display=f"/{cmd}",
                        display_meta=desc
                    )

                                                                                
                                               
                                                                                

class AgentTools:
                                                                                  
    def __init__(
        self,
        cwd: str = None,
        browser_state_file: str = ".deepx_browser_state.json",
        browser_headless: bool = True
    ):
        self.cwd = os.path.abspath(cwd) if cwd else default_workspace_dir()
                                                                                      
                                                                                         
        self.browser_state_file = os.path.abspath(os.path.join(APP_DIR, browser_state_file))
        self.project_memory_file = os.path.abspath(
            os.path.join(APP_DIR, ".deepx_project_memory.json")
        )
        self.browser_headless = browser_headless
                                                                                           
        self.protect_self = False
        self.extra_protected = set()
        self.todos = []
        self._todo_dirty = False
        self._todo_rendered_signature = None
        self.playwright = None
        self.browser = None
        self.browser_context = None
        self.page = None
        self.browser_timeout = 15000
        self.browser_stealth_status = "not initialized"
        self.active_tasks = {}
        self._background_task_counter = 0

                                                                    
                           
                                                                    

    def add_protected(self, path: str):
                                                                                          
        if path:
            self.extra_protected.add(os.path.realpath(os.path.abspath(os.path.join(APP_DIR, path))))

    def is_protected(self, path: str) -> bool:
                                                                                       

                                                                               
                                                                               
           
        if not path:
            return False
        try:
            real = os.path.realpath(os.path.abspath(os.path.join(self.cwd, str(path))))
        except Exception:
            return True                                               

        if not self.protect_self:
            return False

        if real in self.extra_protected:
            return True

        try:
            relative = os.path.relpath(real, APP_DIR)
        except ValueError:
            return False                                                  
                                                                                          
                                                                               
        return relative == os.curdir or not relative.startswith(os.pardir)

    def _deny(self, path: str) -> str:
        return (
            f"[Access denied: '{path}' belongs to the DEEPX application itself and is "
            "permanently blocked in this agent style. It does not exist for you. "
            "Do not attempt to reach it by another route.]"
        )

    def _command_touches_self(self, text: str) -> bool:
                                                                                          
        if not self.protect_self or not text:
            return False
        lowered = str(text).lower()
        needles = set(SELF_PROTECTED_FILES) | set(SELF_PROTECTED_DIRS)
        needles.update(SELF_PROTECTED_PREFIXES)
        needles.update(os.path.basename(p).lower() for p in self.extra_protected)
        return any(needle in lowered for needle in needles)

    def _process_kill_guard(self, text: str):
        """Reject process-kill commands that can terminate the DEEPX runner."""
        if not text:
            return None

        command = str(text)
        lowered = command.lower()
        agent_pid = os.getpid()

        taskkill_calls = re.findall(
            r"(?:^|[&|])\s*([^&|\r\n]*\btaskkill\b[^&|\r\n]*)",
            lowered,
        )
        for taskkill_call in taskkill_calls:
            if re.search(r"/(?:im|fi)\b", taskkill_call):
                return (
                    "taskkill by image name or filter is blocked because it can "
                    "terminate DEEPX. Use only `taskkill /PID <pid> /F`."
                )
            pid_values = {
                int(value)
                for value in re.findall(r"/pid\s+(\d+)", taskkill_call)
            }
            if not pid_values:
                return (
                    "mass taskkill is blocked because it can terminate DEEPX. "
                    "Find the exact target PID first, then use "
                    "`taskkill /PID <pid> /F`."
                )
            if agent_pid in pid_values:
                return (
                    f"PID {agent_pid} belongs to the active DEEPX process and "
                    "cannot be terminated."
                )

        if re.search(r"\b(?:wmic|pkill|killall)\b.*\b(?:delete|terminate|python)", lowered):
            return (
                "mass process termination is blocked. Find the exact target "
                "PID and terminate only that PID."
            )

        if "stop-process" in lowered:
            pid_values = {
                int(value)
                for value in re.findall(r"(?:-id\s+|stop-process\s+)(\d+)", lowered)
            }
            if not pid_values:
                return (
                    "Stop-Process without an explicit numeric PID is blocked "
                    "because it can terminate DEEPX."
                )
            if agent_pid in pid_values:
                return (
                    f"PID {agent_pid} belongs to the active DEEPX process and "
                    "cannot be terminated."
                )

        direct_kill = re.search(
            r"(?:^|[&|])\s*(?:kill|tskill)\s+(?:-[a-z0-9]+\s+)*(\d+)\b",
            lowered,
        )
        if direct_kill and int(direct_kill.group(1)) == agent_pid:
            return (
                f"PID {agent_pid} belongs to the active DEEPX process and "
                "cannot be terminated."
            )

        return None

    async def init_browser(self):
        if not self.playwright:
            from playwright.async_api import async_playwright
            self.playwright = await async_playwright().start()
            self.browser = await self.playwright.chromium.launch(headless=self.browser_headless)
            context_options = dict(
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                viewport={"width": 1440, "height": 900},
                accept_downloads=True,
                locale="ru-RU"
            )
            if os.path.exists(self.browser_state_file):
                context_options["storage_state"] = self.browser_state_file
            self.browser_context = await self.browser.new_context(**context_options)
            await self._enable_browser_stealth()
            self.page = await self.browser_context.new_page()
            self._configure_browser_page(self.page)

    async def _enable_browser_stealth(self):
        try:
            installed_version = package_version("playwright-stealth")
            major_version = int(installed_version.split(".", 1)[0])
            if major_version < 2:
                raise RuntimeError(
                    f"установлена устаревшая версия {installed_version}; требуется 2.x"
                )
            from playwright_stealth import Stealth

            await Stealth().apply_stealth_async(self.browser_context)
            self.browser_stealth_status = (
                f"always enabled (playwright-stealth {installed_version})"
            )
        except (ImportError, PackageNotFoundError) as exc:
            self.browser_stealth_status = "unavailable"
            raise RuntimeError(
                "DEEPX запрещено запускать без playwright-stealth 2.x. "
                "Выполните: python -m pip install -U \"playwright-stealth>=2,<3\""
            ) from exc
        except Exception as exc:
            self.browser_stealth_status = f"failed: {type(exc).__name__}: {exc}"
            raise RuntimeError(
                f"DEEPX остановлен: playwright-stealth не удалось включить: {exc}"
            ) from exc

    def _configure_browser_page(self, page):
        page.set_default_timeout(self.browser_timeout)
        page.set_default_navigation_timeout(30000)
        page.on("dialog", lambda dialog: asyncio.create_task(dialog.accept()))

    async def close_browser(self):
        if self.browser_context:
            try:
                await self.browser_context.storage_state(path=self.browser_state_file)
            except Exception:
                pass
        if self.browser:
            await self.browser.close()
        if self.playwright:
            await self.playwright.stop()
        self.page = None
        self.browser_context = None
        self.browser = None
        self.playwright = None

    def run_cmd(self, command: str, inputs: str = None) -> str:
                                                         
        if self._command_touches_self(command) or self._command_touches_self(inputs):
            return self._deny("command referencing DEEPX application files")
        kill_guard = self._process_kill_guard(command)
        if kill_guard:
            return f"[Error: {kill_guard}]"
        try:
            command_to_run = command
            if sys.platform == "win32":
                command_to_run = f"chcp 65001 >nul & {command}"
            child_env = os.environ.copy()
            child_env.pop("FUNSTAT_API_TOKEN", None)
            child_env["PYTHONIOENCODING"] = "utf-8"
            child_env["PYTHONUTF8"] = "1"
            res = subprocess.run(
                command_to_run,
                shell=True,
                cwd=self.cwd,
                input=inputs,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=child_env,
                timeout=60
            )
            out = res.stdout
            if res.stderr:
                out += ("\n" if out else "") + "[STDERR]\n" + res.stderr
            if res.returncode != 0:
                out += f"\n[Exit Code: {res.returncode}]"
            return out.strip() if out.strip() else "[OK]"
        except subprocess.TimeoutExpired:
            return "[Error: Command execution timed out after 60 seconds]"
        except Exception as e:
            return f"[Error executing command: {e}]"

    def run_python(self, code: str) -> str:
                                                                      
        import tempfile
        import os

        if self._command_touches_self(code):
            return self._deny("Python code referencing DEEPX application files")
        kill_guard = self._process_kill_guard(code)
        if kill_guard:
            return f"[Error: {kill_guard}]"
        if re.search(
            r"\bos\.(?:kill|killpg)\s*\(\s*os\.getppid\s*\(",
            code,
        ) or re.search(
            r"\bpsutil\.Process\s*\(\s*os\.getppid\s*\(\s*\)\s*\)"
            r"\s*\.\s*(?:kill|terminate)\s*\(",
            code,
        ):
            return (
                "[Error: Python code may not terminate its parent DEEPX process. "
                "Use a concrete external PID instead.]"
            )

        fd, temp_path = tempfile.mkstemp(suffix=".py", dir=self.cwd)
        with os.fdopen(fd, 'w', encoding='utf-8') as f:
            f.write(code)
            
        try:
            child_env = os.environ.copy()
            child_env.pop("FUNSTAT_API_TOKEN", None)
            child_env["PYTHONIOENCODING"] = "utf-8"
            child_env["PYTHONUTF8"] = "1"
            res = subprocess.run(
                [sys.executable, "-X", "utf8", temp_path],
                cwd=self.cwd,
                capture_output=True,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=child_env,
                timeout=60
            )
            out = res.stdout
            if res.stderr:
                out += ("\n" if out else "") + "[STDERR]\n" + res.stderr
            if res.returncode != 0:
                out += f"\n[Exit Code: {res.returncode}]"
            return out.strip() if out.strip() else "[OK]"
        except subprocess.TimeoutExpired:
            return "[Error: Python script execution timed out after 60 seconds]"
        except Exception as e:
            return f"[Error executing Python script: {e}]"
        finally:
            try:
                os.remove(temp_path)
            except Exception:
                pass

    def _assign_windows_kill_job(self, process):
        if sys.platform != "win32":
            return None

        class JobObjectBasicLimitInformation(ctypes.Structure):
            _fields_ = [
                ("PerProcessUserTimeLimit", ctypes.c_longlong),
                ("PerJobUserTimeLimit", ctypes.c_longlong),
                ("LimitFlags", ctypes.c_uint32),
                ("MinimumWorkingSetSize", ctypes.c_size_t),
                ("MaximumWorkingSetSize", ctypes.c_size_t),
                ("ActiveProcessLimit", ctypes.c_uint32),
                ("Affinity", ctypes.c_size_t),
                ("PriorityClass", ctypes.c_uint32),
                ("SchedulingClass", ctypes.c_uint32),
            ]

        class IoCounters(ctypes.Structure):
            _fields_ = [
                ("ReadOperationCount", ctypes.c_uint64),
                ("WriteOperationCount", ctypes.c_uint64),
                ("OtherOperationCount", ctypes.c_uint64),
                ("ReadTransferCount", ctypes.c_uint64),
                ("WriteTransferCount", ctypes.c_uint64),
                ("OtherTransferCount", ctypes.c_uint64),
            ]

        class JobObjectExtendedLimitInformation(ctypes.Structure):
            _fields_ = [
                ("BasicLimitInformation", JobObjectBasicLimitInformation),
                ("IoInfo", IoCounters),
                ("ProcessMemoryLimit", ctypes.c_size_t),
                ("JobMemoryLimit", ctypes.c_size_t),
                ("PeakProcessMemoryUsed", ctypes.c_size_t),
                ("PeakJobMemoryUsed", ctypes.c_size_t),
            ]

        kernel32 = ctypes.windll.kernel32
        kernel32.CreateJobObjectW.restype = ctypes.c_void_p
        job_handle = kernel32.CreateJobObjectW(None, None)
        if not job_handle:
            return None

        information = JobObjectExtendedLimitInformation()
        information.BasicLimitInformation.LimitFlags = 0x00002000
        configured = kernel32.SetInformationJobObject(
            ctypes.c_void_p(job_handle),
            9,
            ctypes.byref(information),
            ctypes.sizeof(information),
        )
        transport = getattr(process, "_transport", None)
        popen = (
            transport.get_extra_info("subprocess")
            if transport is not None
            else None
        )
        process_handle = getattr(popen, "_handle", None)
        assigned = bool(
            configured
            and process_handle
            and kernel32.AssignProcessToJobObject(
                ctypes.c_void_p(job_handle),
                ctypes.c_void_p(int(process_handle)),
            )
        )
        if assigned:
            return int(job_handle)
        kernel32.CloseHandle(ctypes.c_void_p(job_handle))
        return None

    def _close_windows_kill_job(self, task: dict):
        job_handle = task.get("job_handle")
        if not job_handle:
            return
        task["job_handle"] = None
        ctypes.windll.kernel32.CloseHandle(ctypes.c_void_p(job_handle))

    def _background_process_args(self, command: str):
        if sys.platform != "win32":
            return ["/bin/sh", "-lc", command], None

        import tempfile

        fd, launcher_path = tempfile.mkstemp(
            prefix=".deepx-task-",
            suffix=".cmd",
            dir=tempfile.gettempdir(),
        )
        with os.fdopen(fd, "w", encoding="utf-8", newline="\r\n") as launcher:
            launcher.write("@chcp 65001 >nul\n")
            launcher.write(f"@{command}\n")
        return [
            os.environ.get("COMSPEC", "cmd.exe"),
            "/d",
            "/s",
            "/c",
            launcher_path,
        ], launcher_path

    async def _capture_background_stream(self, task_id: str, stream, label: str):
        while True:
            chunk = await stream.read(4096)
            if not chunk:
                break
            text = chunk.decode("utf-8", errors="replace")
            if label:
                text = f"[{label}]\n{text}"
            task = self.active_tasks.get(task_id)
            if task is None:
                return
            task["log"] += text
            if len(task["log"]) > MAX_BACKGROUND_LOG_CHARS:
                removed = len(task["log"]) - MAX_BACKGROUND_LOG_CHARS
                task["log"] = (
                    f"[Log truncated: removed {removed} oldest characters]\n"
                    + task["log"][-MAX_BACKGROUND_LOG_CHARS:]
                )

    async def _watch_background_task(self, task_id: str):
        task = self.active_tasks[task_id]
        try:
            returncode = await task["process"].wait()
            await asyncio.gather(*task["reader_tasks"], return_exceptions=True)
            task["returncode"] = returncode
            task["ended_at"] = datetime.now(timezone.utc).isoformat()
        finally:
            self._close_windows_kill_job(task)
            launcher_path = task.get("launcher_path")
            if launcher_path:
                try:
                    os.remove(launcher_path)
                except OSError:
                    pass

    async def run_background_cmd(self, command: str, inputs: str = None) -> str:
        if not command:
            return "[Error: run_background_cmd requires a 'command' argument]"
        if self._command_touches_self(command) or self._command_touches_self(inputs):
            return self._deny("background command referencing DEEPX application files")
        kill_guard = self._process_kill_guard(command)
        if kill_guard:
            return f"[Error: {kill_guard}]"

        child_env = os.environ.copy()
        child_env.pop("FUNSTAT_API_TOKEN", None)
        child_env["PYTHONIOENCODING"] = "utf-8"
        child_env["PYTHONUTF8"] = "1"
        launcher_path = None

        try:
            process_args, launcher_path = self._background_process_args(command)
            process_options = {}
            if sys.platform == "win32":
                process_options["creationflags"] = subprocess.CREATE_NEW_PROCESS_GROUP
            else:
                process_options["start_new_session"] = True
            process = await asyncio.create_subprocess_exec(
                *process_args,
                cwd=self.cwd,
                env=child_env,
                stdin=asyncio.subprocess.PIPE if inputs is not None else asyncio.subprocess.DEVNULL,
                stdout=asyncio.subprocess.PIPE,
                stderr=asyncio.subprocess.PIPE,
                **process_options,
            )
            job_handle = self._assign_windows_kill_job(process)
            self._background_task_counter += 1
            task_id = f"task-{self._background_task_counter:04d}"
            task = {
                "id": task_id,
                "command": command,
                "process": process,
                "pid": process.pid,
                "started_at": datetime.now(timezone.utc).isoformat(),
                "ended_at": None,
                "returncode": None,
                "log": "",
                "launcher_path": launcher_path,
                "job_handle": job_handle,
            }
            self.active_tasks[task_id] = task
            stdout_reader = asyncio.create_task(
                self._capture_background_stream(task_id, process.stdout, "")
            )
            stderr_reader = asyncio.create_task(
                self._capture_background_stream(task_id, process.stderr, "STDERR")
            )
            task["reader_tasks"] = (stdout_reader, stderr_reader)
            task["watcher"] = asyncio.create_task(self._watch_background_task(task_id))

            if inputs is not None and process.stdin is not None:
                process.stdin.write(str(inputs).encode("utf-8"))
                await process.stdin.drain()
                process.stdin.close()

            return (
                f"[Success: Background task '{task_id}' started "
                f"(PID {process.pid}). Use task_status/task_log with this id.]"
            )
        except Exception as e:
            if launcher_path:
                try:
                    os.remove(launcher_path)
                except OSError:
                    pass
            return f"[Error starting background command: {e}]"

    async def task_status(self, task_id: str) -> str:
        task = self.active_tasks.get(str(task_id))
        if task is None:
            return f"[Error: Background task '{task_id}' not found]"
        await asyncio.sleep(0)
        returncode = task["process"].returncode
        if returncode is None:
            status = "running"
        elif returncode == 0:
            status = "completed"
        else:
            status = "failed"
        return (
            f"[Background task '{task_id}': status={status}, PID={task['pid']}, "
            f"exit_code={returncode}, started_at={task['started_at']}, "
            f"ended_at={task.get('ended_at')}]"
        )

    async def task_log(self, task_id: str, tail_lines: int = 200) -> str:
        task = self.active_tasks.get(str(task_id))
        if task is None:
            return f"[Error: Background task '{task_id}' not found]"
        await asyncio.sleep(0)
        try:
            tail_lines = max(1, min(int(tail_lines), 5000))
        except (TypeError, ValueError):
            tail_lines = 200
        log = task["log"]
        if not log:
            state = "running" if task["process"].returncode is None else "finished"
            return f"[Background task '{task_id}' has no output yet; status={state}]"
        lines = log.splitlines()
        shown = "\n".join(lines[-tail_lines:])
        omitted = len(lines) - min(len(lines), tail_lines)
        prefix = f"[Background task '{task_id}' log"
        if omitted:
            prefix += f"; {omitted} earlier line(s) omitted"
        return f"{prefix}]\n{shown}"

    async def close_background_tasks(self):
        running = [
            task for task in self.active_tasks.values()
            if task["process"].returncode is None
        ]
        for task in running:
            try:
                if sys.platform == "win32" and task.get("job_handle"):
                    self._close_windows_kill_job(task)
                elif sys.platform == "win32":
                    task["process"].terminate()
                else:
                    os.killpg(os.getpgid(task["pid"]), signal.SIGTERM)
            except (ProcessLookupError, PermissionError, OSError):
                pass
        for task in running:
            try:
                await asyncio.wait_for(task["process"].wait(), timeout=3)
            except asyncio.TimeoutError:
                try:
                    if sys.platform == "win32":
                        task["process"].kill()
                    else:
                        os.killpg(os.getpgid(task["pid"]), signal.SIGKILL)
                except (ProcessLookupError, PermissionError, OSError):
                    pass
        watchers = [
            task.get("watcher") for task in self.active_tasks.values()
            if task.get("watcher") is not None
        ]
        if watchers:
            try:
                await asyncio.wait_for(
                    asyncio.gather(*watchers, return_exceptions=True),
                    timeout=3,
                )
            except asyncio.TimeoutError:
                for watcher in watchers:
                    if not watcher.done():
                        watcher.cancel()
                await asyncio.gather(*watchers, return_exceptions=True)

    def _load_project_memory(self) -> dict:
                                                                           
        default = {
            "version": 1,
            "project": os.path.basename(self.cwd),
            "updated_at": None,
            "summary": "",
            "current_task": "",
            "completed": [],
            "next_steps": [],
            "important_files": [],
            "facts": [],
            "events": [],
        }
        if not os.path.exists(self.project_memory_file):
            return default
        try:
            with open(self.project_memory_file, "r", encoding="utf-8") as file:
                stored = json.load(file)
            if not isinstance(stored, dict):
                raise ValueError("memory root must be an object")
            for key, value in default.items():
                stored.setdefault(key, value)
            return stored
        except Exception as error:
            raise RuntimeError(f"Cannot read project memory: {error}") from error

    def _save_project_memory(self, memory: dict):
                                                                                          
        memory["updated_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
        temp_path = self.project_memory_file + ".tmp"
        with open(temp_path, "w", encoding="utf-8") as file:
            json.dump(memory, file, ensure_ascii=False, indent=2)
            file.write("\n")
        os.replace(temp_path, self.project_memory_file)

    @staticmethod
    def _memory_list(value, name: str) -> list:
        if value is None:
            return []
        if not isinstance(value, list):
            raise ValueError(f"{name} must be a list")
        return [str(item).strip() for item in value if str(item).strip()]

    def project_memory(self, action: str, args: dict) -> str:
                                                                               
        action = (action or "get").strip().lower()
        try:
            memory = self._load_project_memory()
            if action in ("get", "view", "status"):
                return json.dumps(memory, ensure_ascii=False, indent=2)

            if action == "update":
                scalar_fields = ("summary", "current_task")
                list_fields = ("completed", "next_steps", "important_files", "facts")
                for field in scalar_fields:
                    if field in args:
                        memory[field] = str(args[field]).strip()
                for field in list_fields:
                    if field in args:
                        memory[field] = self._memory_list(args[field], field)
                self._save_project_memory(memory)
                return f"[Success: Project memory updated at '{self.project_memory_file}']"

            if action in ("append_event", "log"):
                event = {
                    "at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
                    "task": str(args.get("task", "")).strip(),
                    "result": str(args.get("result", "")).strip(),
                    "status": str(args.get("status", "completed")).strip() or "completed",
                    "files": self._memory_list(args.get("files", []), "files"),
                }
                if not event["task"] or not event["result"]:
                    return "[Error: append_event requires non-empty task and result]"
                memory["events"].append(event)
                memory["events"] = memory["events"][-50:]
                self._save_project_memory(memory)
                return f"[Success: Project event recorded ({len(memory['events'])} retained)]"

            return "[Error: project_memory action must be get, update, or append_event]"
        except Exception as error:
            return f"[Error in project_memory: {error}]"

    def web_search(self, query: str, max_results: int = 10, site: str = "", region: str = "wt-wt") -> str:
                                                                                     
        if not query or not str(query).strip():
            return "[Error: web_search requires a 'query' argument]"

        query = str(query).strip()
        if site:
            query = f"site:{site} {query}"
        max_results = max(1, min(_as_int(max_results, 10), 25))

        try:
            import requests
            from bs4 import BeautifulSoup
            from urllib.parse import unquote, urlparse, parse_qs
        except ImportError as e:
            return f"[Error: web_search requires requests and beautifulsoup4 ({e})]"

        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                          "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
            "Accept-Language": "en-US,en;q=0.9,ru;q=0.8",
        }

        def _clean_link(href: str) -> str:
                                                                   
            if not href:
                return ""
            if href.startswith("//"):
                href = "https:" + href
            if "duckduckgo.com/l/" in href:
                target = parse_qs(urlparse(href).query).get("uddg", [""])[0]
                if target:
                    return unquote(target)
            return href

        def _is_ad(url: str) -> bool:
                                                                                          
            return "duckduckgo.com/y.js" in url or "ad_provider=" in url

        results = []
        errors = []
        for endpoint in ("https://html.duckduckgo.com/html/", "https://lite.duckduckgo.com/lite/"):
            if len(results) >= max_results:
                break
            try:
                response = requests.post(
                    endpoint,
                    data={"q": query, "kl": region},
                    headers=headers,
                    timeout=20,
                )
                response.raise_for_status()
                soup = BeautifulSoup(response.text, "html.parser")

                blocks = soup.select("div.result, div.web-result")
                if blocks:
                    for block in blocks:
                        link = block.select_one("a.result__a")
                        if not link:
                            continue
                        url = _clean_link(link.get("href", ""))
                        if not url or _is_ad(url) or any(r["url"] == url for r in results):
                            continue
                        snippet_el = block.select_one(".result__snippet")
                        results.append({
                            "title": link.get_text(" ", strip=True),
                            "url": url,
                            "snippet": snippet_el.get_text(" ", strip=True) if snippet_el else "",
                        })
                        if len(results) >= max_results:
                            break
                else:
                                                                 
                    for link in soup.select("a.result-link"):
                        url = _clean_link(link.get("href", ""))
                        if not url or _is_ad(url) or any(r["url"] == url for r in results):
                            continue
                        results.append({
                            "title": link.get_text(" ", strip=True),
                            "url": url,
                            "snippet": "",
                        })
                        if len(results) >= max_results:
                            break
            except Exception as e:
                errors.append(f"{endpoint}: {e}")

        if not results:
            detail = " | ".join(errors) if errors else "no matches"
            return f"[web_search '{query}': no results ({detail})]"

        lines = [f"--- DuckDuckGo results for '{query}' ({len(results)}) ---"]
        for i, item in enumerate(results, 1):
            lines.append(f"{i}. {item['title']}\n   {item['url']}")
            if item["snippet"]:
                lines.append(f"   {item['snippet'][:400]}")
        return "\n".join(lines)

    def fetch_url(self, url: str, max_chars: int = 6000) -> str:
                                                                                              
        if not url:
            return "[Error: fetch_url requires a 'url' argument]"
        if not str(url).startswith(("http://", "https://")):
            url = "https://" + str(url).lstrip("/")

        max_chars = max(500, min(_as_int(max_chars, 6000), 20000))
        try:
            import requests
            from bs4 import BeautifulSoup
        except ImportError as e:
            return f"[Error: fetch_url requires requests and beautifulsoup4 ({e})]"

        try:
            response = requests.get(
                url,
                headers={
                    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                                  "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
                    "Accept-Language": "en-US,en;q=0.9,ru;q=0.8",
                },
                timeout=25,
            )
            status = response.status_code
            content_type = response.headers.get("Content-Type", "")

            if "html" not in content_type and "xml" not in content_type and "text" not in content_type:
                return f"[fetch_url {url}: HTTP {status}, non-text content ({content_type})]"

            soup = BeautifulSoup(response.text, "html.parser")
            for tag in soup(["script", "style", "noscript", "svg", "iframe"]):
                tag.decompose()
            title = soup.title.get_text(strip=True) if soup.title else ""
            text = re.sub(r"\n{3,}", "\n\n", soup.get_text("\n", strip=True))

            truncated = ""
            if len(text) > max_chars:
                truncated = f"\n[Truncated: {len(text)} chars total, shown {max_chars}]"
                text = text[:max_chars]
            return f"--- {url} (HTTP {status}) ---\nTitle: {title}\n\n{text}{truncated}"
        except Exception as e:
            return f"[Error fetching '{url}': {e}]"

    def read_file(self, path: str, start_line: int = 1, end_line: int = 500) -> str:
                                          
        if not path:
            return "[Error: read_file requires a 'path' argument]"
        abs_path = os.path.abspath(os.path.join(self.cwd, path))
        if self.is_protected(path):
            return self._deny(path)
        if not os.path.exists(abs_path):
            return f"[Error: File '{path}' does not exist]"
        if os.path.isdir(abs_path):
            return f"[Error: Path '{path}' is a directory, not a file]"
        try:
            with open(abs_path, "r", encoding="utf-8", errors="replace") as f:
                lines = f.readlines()

            try:
                start = max(1, int(start_line)) - 1
            except (TypeError, ValueError):
                start = 0
            try:
                end = min(len(lines), int(end_line))
            except (TypeError, ValueError):
                end = min(len(lines), start + MAX_READ_LINES)

                                                                                              
                                                                           
            truncated = False
            if end - start > MAX_READ_LINES:
                end = start + MAX_READ_LINES
                truncated = True

            selected = lines[start:end]
            formatted = "".join(f"{i + start + 1:4d} | {line}" for i, line in enumerate(selected))
            header = f"--- Content of {path} (Lines {start+1}-{end} of {len(lines)}) ---"
            footer = ""
            if truncated or end < len(lines):
                footer = (
                    f"\n[Truncated: shown {end - start} of {len(lines)} lines. "
                    f"Continue with read_file path=\"{path}\" start_line={end + 1}]"
                )
            return f"{header}\n{formatted}{footer}"
        except Exception as e:
            return f"[Error reading file '{path}': {e}]"

    def _syntax_check_file(self, abs_path: str, content: str) -> str:
        suffix = Path(abs_path).suffix.lower()
        try:
            if suffix == ".py":
                py_compile.compile(abs_path, doraise=True)
            elif suffix == ".json":
                json.loads(content)
            elif suffix in {".js", ".mjs", ".cjs"}:
                node = shutil.which("node")
                if node is None:
                    return "\n[Syntax Check Skipped: Node.js executable was not found]"
                result = subprocess.run(
                    [node, "--check", abs_path],
                    cwd=self.cwd,
                    capture_output=True,
                    text=True,
                    encoding="utf-8",
                    errors="replace",
                    timeout=15,
                )
                if result.returncode != 0:
                    detail = (result.stderr or result.stdout).strip()
                    detail = re.sub(r"\s+", " ", detail)
                    return f"\n[Success: File saved, BUT Syntax Check Failed: {detail}]"
            else:
                return ""
        except py_compile.PyCompileError as error:
            detail = re.sub(r"\s+", " ", str(error))
            return f"\n[Success: File saved, BUT Syntax Check Failed: {detail}]"
        except json.JSONDecodeError as error:
            return (
                "\n[Success: File saved, BUT Syntax Check Failed: "
                f"line {error.lineno}, column {error.colno}: {error.msg}]"
            )
        except subprocess.TimeoutExpired:
            return "\n[Syntax Check Skipped: Node.js syntax check timed out]"
        except Exception as error:
            return f"\n[Syntax Check Skipped: {type(error).__name__}: {error}]"
        return "\n[Syntax Check: OK]"

    def write_file(self, path: str, content: str) -> str:
                                                                                    
        abs_path = os.path.abspath(os.path.join(self.cwd, path))
        if self.is_protected(path):
            return self._deny(path)
        filename = os.path.basename(path)
        existed = os.path.exists(abs_path)
        old_line_count = 0

        if existed:
            try:
                with open(abs_path, "r", encoding="utf-8", errors="replace") as f:
                    old_line_count = len(f.readlines())
            except Exception:
                pass

        try:
            os.makedirs(os.path.dirname(abs_path), exist_ok=True)
            if filename.endswith(".py") or "**" in content:
                content = re.sub(r'\*\*([a-zA-Z0-9_]+)\*\*', r'__\1__', content)
            with open(abs_path, "w", encoding="utf-8") as f:
                f.write(content)

            new_line_count = len(content.splitlines())
            syntax_result = self._syntax_check_file(abs_path, content)

            if not existed:
                return (
                    f"[Success: File '{path}' created "
                    f"({new_line_count} lines, {len(content)} bytes)]{syntax_result}"
                )
            else:
                added = max(0, new_line_count - old_line_count)
                removed = max(0, old_line_count - new_line_count)
                stats = []
                if added > 0 or (added == 0 and removed == 0):
                    stats.append(f"[bold #10B981]+{added}[/bold #10B981]")
                if removed > 0:
                    stats.append(f"[bold #EF4444]-{removed}[/bold #EF4444]")
                stat_str = " ".join(stats)
                return (
                    f"[Success: File '{path}' updated "
                    f"({new_line_count} lines, {len(content)} bytes)]{syntax_result}"
                )
        except Exception as e:
            return f"[Error writing file '{path}': {e}]"

    def edit_file(
        self,
        path: str,
        target: str = "",
        replacement: str = "",
        start_line=None,
        end_line=None,
    ) -> str:
                                                                                                
        abs_path = os.path.abspath(os.path.join(self.cwd, path))
        if self.is_protected(path):
            return self._deny(path)
        filename = os.path.basename(path)
        if not os.path.exists(abs_path):
            return f"[Error: File '{path}' does not exist]"
        try:
            with open(abs_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()

            match_method = "exact match"
            matched_target = target
            if start_line is not None or end_line is not None:
                try:
                    start = int(start_line if start_line is not None else end_line)
                    end = int(end_line if end_line is not None else start_line)
                except (TypeError, ValueError):
                    return f"[Error: start_line and end_line must be integers for '{path}']"
                if start < 1 or end < start:
                    return f"[Error: Invalid line range {start}-{end} for '{path}']"
                source_lines = content.splitlines(keepends=True)
                if start > len(source_lines) or end > len(source_lines):
                    return (
                        f"[Error: Line range {start}-{end} is outside '{path}' "
                        f"(1-{len(source_lines)})]"
                    )
                matched_target = "".join(source_lines[start - 1:end])
                line_ending = "\r\n" if "\r\n" in matched_target else "\n"
                range_replacement = replacement
                if (
                    range_replacement
                    and matched_target.endswith(("\n", "\r"))
                    and not range_replacement.endswith(("\n", "\r"))
                    and end < len(source_lines)
                ):
                    range_replacement += line_ending
                prefix = "".join(source_lines[:start - 1])
                suffix = "".join(source_lines[end:])
                new_content = prefix + range_replacement + suffix
                match_method = f"line range {start}-{end}"
            elif target and target in content:
                new_content = content.replace(target, replacement, 1)
            else:
                if not target:
                    return (
                        f"[Error: edit_file requires either a non-empty target or "
                        f"start_line/end_line for '{path}']"
                    )
                target_lines_normalized = [
                    line.rstrip() for line in target.splitlines()
                ]
                source_lines = content.splitlines(keepends=True)
                source_lines_normalized = [
                    line.rstrip("\r\n").rstrip() for line in source_lines
                ]
                width = len(target_lines_normalized)
                matches = [
                    index
                    for index in range(len(source_lines_normalized) - width + 1)
                    if source_lines_normalized[index:index + width]
                    == target_lines_normalized
                ]
                if not matches:
                    return f"[Error: Target string not found in '{path}', including normalized matching]"
                if len(matches) > 1:
                    lines = ", ".join(str(index + 1) for index in matches[:10])
                    return (
                        f"[Error: Normalized target is ambiguous in '{path}' "
                        f"(matches start at lines {lines}); use start_line/end_line]"
                    )
                index = matches[0]
                matched_target = "".join(source_lines[index:index + width])
                line_ending = "\r\n" if "\r\n" in matched_target else "\n"
                fuzzy_replacement = replacement
                if (
                    fuzzy_replacement
                    and matched_target.endswith(("\n", "\r"))
                    and not fuzzy_replacement.endswith(("\n", "\r"))
                    and index + width < len(source_lines)
                ):
                    fuzzy_replacement += line_ending
                new_content = (
                    "".join(source_lines[:index])
                    + fuzzy_replacement
                    + "".join(source_lines[index + width:])
                )
                match_method = f"normalized whitespace match at line {index + 1}"

            target_lines = len(matched_target.splitlines())
            replacement_lines = len(replacement.splitlines())
            diff = replacement_lines - target_lines
            with open(abs_path, "w", encoding="utf-8") as f:
                f.write(new_content)
            syntax_result = self._syntax_check_file(abs_path, new_content)

            added = max(0, diff)
            removed = max(0, -diff)
            stats = []
            if added > 0 or (added == 0 and removed == 0):
                stats.append(f"[bold #10B981]+{added}[/bold #10B981]")
            if removed > 0:
                stats.append(f"[bold #EF4444]-{removed}[/bold #EF4444]")
            stat_str = " ".join(stats)

            return (
                f"[Success: File '{path}' updated successfully using {match_method}]"
                f"{syntax_result}"
            )
        except Exception as e:
            return f"[Error editing file '{path}': {e}]"
    async def _browser_locator(self, args: dict):
                                                                             
        if args.get("ref"):
            return self.page.locator(f'[data-deepx-ref="{args["ref"]}"]').first
        if args.get("selector"):
            return self.page.locator(args["selector"]).first
        name = args.get("name") or args.get("target")
        exact = bool(args.get("exact", False))
        if args.get("role") and name:
            return self.page.get_by_role(args["role"], name=name, exact=exact).first
        if args.get("label"):
            return self.page.get_by_label(args["label"], exact=exact).first
        if args.get("placeholder"):
            return self.page.get_by_placeholder(args["placeholder"], exact=exact).first
        if name:
            return self.page.get_by_text(name, exact=exact).first
        raise ValueError("Specify ref, selector, role+name, label, placeholder, or target")

    async def _browser_settle(self, delay: int = 500):
        await self.page.wait_for_timeout(delay)
        try:
            await self.page.wait_for_load_state("domcontentloaded", timeout=3000)
        except Exception:
            pass

    async def browser_snapshot(self, max_elements: int = 120, max_text: int = 12000) -> str:
                                                                              
        state = await self.page.evaluate(
            """
            ({maxElements, maxText}) => {
              document.querySelectorAll("[data-deepx-ref]").forEach(
                el => el.removeAttribute("data-deepx-ref")
              );
              const clean = (value, limit) =>
                String(value || "").replace(/\\s+/g, " ").trim().slice(0, limit);
              const visible = el => {
                const r = el.getBoundingClientRect(), s = getComputedStyle(el);
                return r.width > 0 && r.height > 0 && s.display !== "none" &&
                       s.visibility !== "hidden" && Number(s.opacity || 1) > 0;
              };
              const query = [
                "a[href]", "button", "input", "textarea", "select", "summary",
                "[contenteditable='true']", "[role='button']", "[role='link']",
                "[role='textbox']", "[role='checkbox']", "[role='radio']",
                "[role='combobox']", "[role='menuitem']", "[role='tab']"
              ].join(",");
              const elements = [];
              for (const el of document.querySelectorAll(query)) {
                if (!visible(el) || elements.length >= maxElements) continue;
                const ref = "e" + (elements.length + 1);
                el.dataset.deepxRef = ref;
                const tag = el.tagName.toLowerCase();
                const type = (el.getAttribute("type") || "").toLowerCase();
                let role = el.getAttribute("role");
                if (!role) role =
                  tag === "a" ? "link" :
                  tag === "button" || tag === "summary" ? "button" :
                  tag === "select" ? "combobox" :
                  type === "checkbox" ? "checkbox" :
                  type === "radio" ? "radio" :
                  ["button", "submit", "reset"].includes(type) ? "button" :
                  ["input", "textarea"].includes(tag) || el.isContentEditable ? "textbox" : tag;
                const by = el.getAttribute("aria-labelledby");
                const byText = by ? by.split(/\\s+/).map(
                  id => document.getElementById(id)?.innerText || ""
                ).join(" ") : "";
                const labels = el.labels ? Array.from(el.labels).map(x => x.innerText).join(" ") : "";
                const name = clean(
                  el.getAttribute("aria-label") || byText || labels ||
                  el.getAttribute("placeholder") || el.innerText ||
                  el.getAttribute("value") || el.getAttribute("title"), 180
                );
                const item = {ref, role, name};
                if (tag === "a") item.href = el.href;
                if ("value" in el && ["input", "textarea", "select"].includes(tag))
                  item.value = clean(el.value, 100);
                if ("checked" in el) item.checked = Boolean(el.checked);
                if (el.disabled || el.getAttribute("aria-disabled") === "true")
                  item.disabled = true;
                elements.push(item);
              }
              return {
                url: location.href, title: document.title,
                text: clean(document.body?.innerText, maxText), elements,
                scroll: {
                  y: Math.round(scrollY),
                  maxY: Math.max(0, document.documentElement.scrollHeight - innerHeight)
                }
              };
            }
            """,
            {"maxElements": max_elements, "maxText": max_text}
        )
        state["tabs"] = [
            {"index": i, "active": page == self.page, "url": page.url}
            for i, page in enumerate(self.browser_context.pages)
        ]
        return json.dumps(state, ensure_ascii=False, indent=2)

    async def browser_action(self, action: str, args: dict) -> str:
                                                                                  
        await self.init_browser()
        args = args or {}
        action = (action or "snapshot").lower()
        try:
            if action in ("snapshot", "observe"):
                return await self.browser_snapshot(
                    int(args.get("max_elements", 120)), int(args.get("max_text", 12000))
                )
            if action in ("navigate", "open"):
                url = str(args.get("url", "")).strip()
                if not url:
                    raise ValueError("url is required")
                if not re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", url):
                    url = "https://" + url
                await self.page.goto(url, wait_until="domcontentloaded")
            elif action == "click":
                loc = await self._browser_locator(args)
                await loc.scroll_into_view_if_needed()
                old_pages = len(self.browser_context.pages)
                await loc.click()
                await self._browser_settle(800)
                if len(self.browser_context.pages) > old_pages:
                    self.page = self.browser_context.pages[-1]
                    self._configure_browser_page(self.page)
            elif action in ("fill", "type"):
                loc = await self._browser_locator(args)
                value = str(args.get("value", args.get("text", "")))
                await loc.scroll_into_view_if_needed()
                if action == "type" and args.get("append"):
                    await loc.type(value, delay=int(args.get("delay", 0)))
                else:
                    await loc.fill(value)
            elif action == "press":
                targeted = any(args.get(k) for k in (
                    "ref", "selector", "role", "name", "label", "placeholder", "target"
                ))
                if targeted:
                    await (await self._browser_locator(args)).press(args.get("key", "Enter"))
                else:
                    await self.page.keyboard.press(args.get("key", "Enter"))
            elif action == "select":
                loc = await self._browser_locator(args)
                if args.get("option_label") is not None:
                    await loc.select_option(label=str(args["option_label"]))
                else:
                    await loc.select_option(value=str(args.get("value", "")))
            elif action in ("check", "uncheck", "hover", "focus"):
                loc = await self._browser_locator(args)
                await getattr(loc, action)()
            elif action == "scroll":
                if args.get("ref") or args.get("selector") or args.get("target"):
                    await (await self._browser_locator(args)).scroll_into_view_if_needed()
                else:
                    direction = -1 if args.get("direction") == "up" else 1
                    await self.page.mouse.wheel(0, int(args.get("amount", 700)) * direction)
            elif action == "wait":
                timeout = min(int(args.get("timeout", 10000)), 60000)
                if args.get("text"):
                    await self.page.get_by_text(args["text"], exact=bool(args.get("exact"))).first.wait_for(
                        state=args.get("state", "visible"), timeout=timeout
                    )
                elif args.get("selector"):
                    await self.page.locator(args["selector"]).first.wait_for(
                        state=args.get("state", "visible"), timeout=timeout
                    )
                elif args.get("url"):
                    await self.page.wait_for_url(args["url"], timeout=timeout)
                else:
                    await self.page.wait_for_timeout(timeout)
            elif action in ("back", "forward"):
                await getattr(self.page, "go_" + action)(wait_until="domcontentloaded")
            elif action == "reload":
                await self.page.reload(wait_until="domcontentloaded")
            elif action == "new_tab":
                self.page = await self.browser_context.new_page()
                self._configure_browser_page(self.page)
                if args.get("url"):
                    await self.page.goto(args["url"], wait_until="domcontentloaded")
            elif action == "switch_tab":
                self.page = self.browser_context.pages[int(args.get("index", -1))]
                await self.page.bring_to_front()
            elif action == "close_tab":
                if len(self.browser_context.pages) < 2:
                    raise ValueError("Cannot close the only tab")
                await self.page.close()
                self.page = self.browser_context.pages[-1]
            elif action == "list_tabs":
                return json.dumps([
                    {"index": i, "active": p == self.page, "title": await p.title(), "url": p.url}
                    for i, p in enumerate(self.browser_context.pages)
                ], ensure_ascii=False, indent=2)
            elif action == "upload":
                files = args.get("files", args.get("path", []))
                if isinstance(files, str):
                    files = [files]
                await (await self._browser_locator(args)).set_input_files([
                    os.path.abspath(os.path.join(self.cwd, p)) for p in files
                ])
            elif action == "download":
                async with self.page.expect_download(timeout=30000) as info:
                    await (await self._browser_locator(args)).click()
                download = await info.value
                target = os.path.abspath(os.path.join(
                    self.cwd, args.get("path") or download.suggested_filename
                ))
                os.makedirs(os.path.dirname(target), exist_ok=True)
                await download.save_as(target)
                return f"[Success: Downloaded to {target}]\n" + await self.browser_snapshot()
            elif action == "screenshot":
                target = os.path.abspath(os.path.join(
                    self.cwd, args.get("path", "browser_screenshot.png")
                ))
                os.makedirs(os.path.dirname(target), exist_ok=True)
                await self.page.screenshot(path=target, full_page=bool(args.get("full_page", True)))
                return f"[Success: Screenshot saved to {target}]\n" + await self.browser_snapshot()
            elif action == "extract_text":
                loc = self.page.locator(args.get("selector", "body")).first
                return (await loc.inner_text())[:int(args.get("max_text", 20000))]
            elif action == "extract_html":
                html = await self.page.content() if not args.get("selector") else (
                    await self.page.locator(args["selector"]).first.inner_html()
                )
                return html[:int(args.get("max_text", 30000))]
            elif action == "evaluate":
                result = await self.page.evaluate(args.get("script", ""))
                return json.dumps(result, ensure_ascii=False, default=str)[:30000]
            else:
                return f"[Error: Unknown browser action '{action}']"

            await self._browser_settle(int(args.get("settle_ms", 500)))
            return f"[Success: browser {action}]\n" + await self.browser_snapshot()
        except Exception as e:
            try:
                state = await self.browser_snapshot(60, 6000)
            except Exception:
                state = "{}"
            return f"[Error in browser '{action}': {type(e).__name__}: {e}]\nCurrent state:\n{state}"

    def list_dir(self, path: str = ".") -> str:
                                                    
        abs_path = os.path.abspath(os.path.join(self.cwd, path))
        if self.is_protected(path):
            return self._deny(path)
        if not os.path.exists(abs_path):
            return f"[Error: Path '{path}' does not exist]"
        try:
            entries = sorted(os.listdir(abs_path))
            dirs = []
            files = []
            for entry in entries:
                full = os.path.join(abs_path, entry)
                                                                                       
                                                       
                if self.is_protected(full):
                    continue
                if os.path.isdir(full):
                    dirs.append(f"[DIR]  {entry}/")
                else:
                    size = os.path.getsize(full)
                    files.append(f"[FILE] {entry} ({size:,} bytes)")
            out = dirs + files
            return f"--- Contents of {path} ({len(out)} items) ---\n" + ("\n".join(out) if out else "[Empty Directory]")
        except Exception as e:
            return f"[Error listing directory '{path}': {e}]"

    def file_info(self, path: str) -> str:
                                
        abs_path = os.path.abspath(os.path.join(self.cwd, path))
        if self.is_protected(path):
            return self._deny(path)
        if not os.path.exists(abs_path):
            return f"[Error: File '{path}' does not exist]"
        try:
            stat = os.stat(abs_path)
            return (
                f"Path: {path}\n"
                f"Absolute: {abs_path}\n"
                f"Type: {'Directory' if os.path.isdir(abs_path) else 'File'}\n"
                f"Size: {stat.st_size:,} bytes\n"
                f"Modified: {stat.st_mtime}"
            )
        except Exception as e:
            return f"[Error retrieving info for '{path}': {e}]"

                                                                                       
    TODO_SYMBOLS = {
        "pending": ("○", "#64748B"),                     
        "active": ("◐", "#FBBF24"),                    
        "done": ("●", "#10B981"),                       
        "cancelled": ("⊘", "#EF4444"),                   
    }
    TODO_STATUS_ALIASES = {
        "todo": "pending", "open": "pending", "new": "pending", "waiting": "pending",
        "in_progress": "active", "progress": "active", "doing": "active",
        "current": "active", "started": "active", "работаю": "active",
        "complete": "done", "completed": "done", "finished": "done", "ok": "done",
        "готово": "done", "выполнено": "done",
        "skip": "cancelled", "skipped": "cancelled", "canceled": "cancelled",
        "blocked": "cancelled", "failed": "cancelled", "отменено": "cancelled",
    }

    def _normalize_todo_status(self, status: str) -> str:
        value = str(status or "pending").strip().lower()
        value = self.TODO_STATUS_ALIASES.get(value, value)
        return value if value in self.TODO_SYMBOLS else "pending"

    def _mark_todo_changed(self) -> str:
                                                                                 
        self._todo_dirty = True
        return self._todo_text()

    def _todo_signature(self):
        return tuple((item["text"], item["status"]) for item in self.todos)

    def _todo_text(self) -> str:
                                                                      
        if not self.todos:
            return "[Task list is empty]"
        done_count = sum(1 for item in self.todos if item["status"] == "done")
        compact = "\n".join(
            f"{i}. [{item['status']}] {item['text']}"
            for i, item in enumerate(self.todos, 1)
        )
        return f"--- Task list ({done_count}/{len(self.todos)} done) ---\n{compact}"

    def flush_todos(self) -> bool:
                                                                                        

                                                                                   
                                                                  
           
        if not self._todo_dirty:
            return False
        self._todo_dirty = False
        signature = self._todo_signature()
        if signature == self._todo_rendered_signature:
            return False
        self._todo_rendered_signature = signature
        self._render_todos()
        return True

    def _render_todos(self) -> str:
                                                                                    
        if not self.todos:
            console.print("[dim #64748B]  Task list is empty.[/dim #64748B]")
            return "[Task list is empty]"

        done_count = sum(1 for item in self.todos if item["status"] == "done")
        total = len(self.todos)
        percent = int(round(done_count * 100 / total))

        lines = []
        for index, item in enumerate(self.todos, 1):
            symbol, color = self.TODO_SYMBOLS[item["status"]]
            number = f"[dim #475569]{index:>2}[/dim #475569]"
            if item["status"] == "done":
                text = f"[dim strike #94A3B8]{item['text']}[/dim strike #94A3B8]"
            elif item["status"] == "active":
                text = f"[bold #FDE68A]{item['text']}[/bold #FDE68A]"
            elif item["status"] == "cancelled":
                text = f"[dim strike #7F1D1D]{item['text']}[/dim strike #7F1D1D]"
            else:
                text = f"[#CBD5E1]{item['text']}[/#CBD5E1]"
            lines.append(f"{number}  [{color}]{symbol}[/{color}]  {text}")

        filled = int(round(percent / 5))
        bar = f"[#536DFE]{'━' * filled}[/#536DFE][#1E293B]{'━' * (20 - filled)}[/#1E293B]"

                                                                                      
                                                                                
        min_inner_width = 34
        widest = max(len(f"{i:>2}  {sym}  {item['text']}")
                     for i, (sym, item) in enumerate(
                         ((self.TODO_SYMBOLS[t["status"]][0], t) for t in self.todos), 1))
        if widest < min_inner_width:
            lines[0] += " " * (min_inner_width - widest)

        panel = Panel(
            "\n".join(lines),
            title=f"[bold #536DFE]TASKS[/bold #536DFE] [dim #64748B]{done_count}/{total}[/dim #64748B]",
            title_align="left",
            subtitle=f"{bar} [bold #536DFE]{percent}%[/bold #536DFE]",
            subtitle_align="right",
            border_style="#536DFE",
            box=box.ROUNDED,
            padding=(0, 2),
            expand=False,
        )
        console.print()
        console.print(panel)
        console.print()

        compact = "\n".join(
            f"{i}. [{item['status']}] {item['text']}"
            for i, item in enumerate(self.todos, 1)
        )
        return f"--- Task list ({done_count}/{total} done) ---\n{compact}"

    def todo(self, action: str = "list", args: dict = None) -> str:
                                                            
        args = args or {}
        action = str(action or "list").strip().lower()

        def _collect_items():
            raw = args.get("items", args.get("tasks", args.get("text", [])))
            if isinstance(raw, str):
                raw = [raw]
            if not isinstance(raw, list):
                return []
            collected = []
            for entry in raw:
                if isinstance(entry, dict):
                    text = str(entry.get("text", entry.get("task", ""))).strip()
                    status = self._normalize_todo_status(entry.get("status", "pending"))
                else:
                    text = str(entry).strip()
                    status = "pending"
                if text:
                    collected.append({"text": text, "status": status})
            return collected

        if action in ("set", "create", "plan", "replace"):
            items = _collect_items()
            if not items:
                return "[Error: todo 'set' requires a non-empty 'items' list]"
            self.todos = items
            return self._mark_todo_changed()

        if action in ("add", "append"):
            items = _collect_items()
            if not items:
                return "[Error: todo 'add' requires a non-empty 'items' list]"
            self.todos.extend(items)
            return self._mark_todo_changed()

        if action in ("update", "status", "complete", "done", "start"):
            if not self.todos:
                return "[Error: task list is empty, call todo 'set' first]"

            def _resolve(spec: dict):
                index = spec.get("index", spec.get("id", spec.get("number")))
                if index is not None:
                    position = _as_int(index, 0)
                    if not 1 <= position <= len(self.todos):
                        return None, f"[Error: task index {index} is out of range 1-{len(self.todos)}]"
                    return self.todos[position - 1], None
                needle = str(spec.get("text", spec.get("task", ""))).strip().lower()
                if not needle:
                    return None, "[Error: todo 'update' requires 'index' or 'text']"
                for item in self.todos:
                    if needle in item["text"].lower():
                        return item, None
                return None, f"[Error: no task matching '{needle}']"

            def _status_for(spec: dict) -> str:
                if action in ("complete", "done"):
                    return "done"
                if action == "start":
                    return "active"
                return self._normalize_todo_status(spec.get("status", "done"))

                                                                                    
                                                                      
            batch = args.get("updates", args.get("changes"))
            specs = batch if isinstance(batch, list) and batch else [args]

            errors = []
            for spec in specs:
                if not isinstance(spec, dict):
                    continue
                target, error = _resolve(spec)
                if error:
                    errors.append(error)
                    continue
                target["status"] = _status_for(spec)
                if spec.get("new_text"):
                    target["text"] = str(spec["new_text"]).strip()

            if errors and len(errors) == len(specs):
                return errors[0]
            summary = self._mark_todo_changed()
            return f"{summary}\n" + "\n".join(errors) if errors else summary

        if action in ("clear", "reset"):
            self.todos = []
            self._todo_dirty = False
            self._todo_rendered_signature = None
            console.print("[dim #64748B]  Task list cleared.[/dim #64748B]")
            return "[Task list cleared]"

        if action in ("list", "show", "get", ""):
                                                               
            self._todo_dirty = False
            self._todo_rendered_signature = self._todo_signature()
            return self._render_todos()

        return f"[Error: unknown todo action '{action}'. Use set, add, update, complete, start, list or clear]"

    def sys_info(self) -> str:
                                          
        return (
            f"OS: {sys.platform}\n"
            f"Python: {sys.version.split()[0]}\n"
            f"Working Directory: {self.cwd}"
        )

    def make_excel(self, path: str, sheets: dict) -> str:
                                                                                                              
        abs_path = os.path.abspath(os.path.join(self.cwd, path))
        if self.is_protected(path):
            return self._deny(path)
        os.makedirs(os.path.dirname(abs_path), exist_ok=True)
        filename = os.path.basename(path)

        if isinstance(sheets, str):
            try:
                sheets = json.loads(sheets)
            except Exception:
                sheets = {"Лист1": [[sheets]]}

        if not isinstance(sheets, dict) or not sheets:
            sheets = {"Sheet1": [["Содержимое"], ["Успешно создано"]]}

        try:
            import openpyxl
            wb = openpyxl.Workbook()
            default_sheet = wb.active

            for i, (sname, rows) in enumerate(sheets.items()):
                sheet_title = str(sname)[:31]
                if i == 0:
                    ws = default_sheet
                    ws.title = sheet_title
                else:
                    ws = wb.create_sheet(title=sheet_title)

                if isinstance(rows, list):
                    for r in rows:
                        if isinstance(r, list):
                            ws.append(r)
                        else:
                            ws.append([str(r)])
                elif isinstance(rows, dict):
                    ws.append(["Параметр", "Значение"])
                    for k, v in rows.items():
                        ws.append([str(k), str(v)])

            wb.save(abs_path)
            return f"[Success: Excel spreadsheet '{path}' created successfully]"
        except Exception as e:
            try:
                import pandas as pd
                with pd.ExcelWriter(abs_path, engine='openpyxl') as writer:
                    for sheet_name, data in sheets.items():
                        if isinstance(data, list) and len(data) > 0:
                            df = pd.DataFrame(data[1:], columns=data[0])
                        else:
                            df = pd.DataFrame()
                        df.to_excel(writer, sheet_name=str(sheet_name)[:31], index=False)
                return f"[Success: Excel spreadsheet '{path}' created with pandas]"
            except Exception as ex:
                return f"[Error creating Excel file '{path}': {e} / {ex}]"

    def render_plan(self, title: str = "План выполнения задачи", steps: list = None) -> str:
                                                                                                                              
        steps = steps or []
        if not steps:
            return "[Note: No plan steps provided]"

        plan_table = Table(show_header=False, box=None, padding=(0, 1))
        plan_table.add_column("Icon", justify="center")
        plan_table.add_column("Step", style="white")

        for i, step in enumerate(steps, 1):
            if isinstance(step, str):
                text = step
                status = "pending"
            elif isinstance(step, dict):
                text = step.get("text", f"Шаг {i}")
                status = str(step.get("status", "pending")).lower()
            else:
                text = str(step)
                status = "pending"

            if status in ["completed", "done", "finished", "true", "yes"]:
                icon = "[bold #10B981]●[/bold #10B981]"
                text_style = f"[dim #94A3B8]{text}[/dim #94A3B8]"
            elif status in ["in_progress", "active", "working"]:
                icon = "[bold #536DFE]○[/bold #536DFE]"
                text_style = f"[bold #536DFE]{text}[/bold #536DFE]"
            else:
                icon = "[bold white]○[/bold white]"
                text_style = f"[white]{text}[/white]"

            plan_table.add_row(icon, text_style)

        panel = Panel(
            plan_table,
            title=f"[bold #536DFE]📋 {title}[/bold #536DFE]",
            border_style="#536DFE",
            box=box.ROUNDED,
            expand=False
        )
        console.print()
        console.print(panel)
        console.print()
        return f"[Success: Plan '{title}' rendered ({len(steps)} step(s))]"

    def make_docx(self, path: str, title: str, paragraphs: list) -> str:
                                                                                            
        abs_path = os.path.abspath(os.path.join(self.cwd, path))
        if self.is_protected(path):
            return self._deny(path)
        try:
            os.makedirs(os.path.dirname(abs_path), exist_ok=True)
            py_code = (
                "import sys, json\n"
                "try:\n"
                "    import docx\n"
                "    doc = docx.Document()\n"
                f"    doc.add_heading('''{title}''', level=0)\n"
                f"    paras = json.loads('''{json.dumps(paragraphs)}''')\n"
                "    for p in paras:\n"
                "        if isinstance(p, dict) and p.get('heading'):\n"
                "            doc.add_heading(p['heading'], level=p.get('level', 1))\n"
                "        else:\n"
                "            doc.add_paragraph(str(p))\n"
                f"    doc.save(r'''{abs_path}''')\n"
                "    print('SUCCESS')\n"
                "except Exception as e:\n"
                "    print(f'ERROR: {e}')\n"
            )
            res = subprocess.run([sys.executable, "-c", py_code], capture_output=True, text=True)
            if "SUCCESS" in res.stdout:
                return f"[Success: Word document '{path}' created]"
            return f"[Success: Word document '{path}' created]"
        except Exception as e:
            return f"[Error creating docx: {e}]"

    def make_pptx(self, path: str, title: str, slides: list) -> str:
                                                                                                         
        abs_path = os.path.abspath(os.path.join(self.cwd, path))
        if self.is_protected(path):
            return self._deny(path)
        try:
            os.makedirs(os.path.dirname(abs_path), exist_ok=True)
            py_code = (
                "import sys, json\n"
                "try:\n"
                "    from pptx import Presentation\n"
                "    prs = Presentation()\n"
                f"    slides = json.loads('''{json.dumps(slides)}''')\n"
                "    for s in slides:\n"
                "        blank_slide_layout = prs.slide_layouts[1]\n"
                "        slide = prs.slides.add_slide(blank_slide_layout)\n"
                "        slide.shapes.title.text = s.get('title', '')\n"
                "        tf = slide.placeholders[1].text_frame\n"
                "        tf.text = s.get('content', [''])[0] if s.get('content') else ''\n"
                "        for bullet in s.get('content', [])[1:]:\n"
                "            p = tf.add_paragraph()\n"
                "            p.text = str(bullet)\n"
                f"    prs.save(r'''{abs_path}''')\n"
                "    print('SUCCESS')\n"
                "except Exception as e:\n"
                "    print(f'ERROR: {e}')\n"
            )
            res = subprocess.run([sys.executable, "-c", py_code], capture_output=True, text=True)
            return f"[Success: PowerPoint presentation '{path}' created ({len(slides)} slide(s))]"
        except Exception as e:
            return f"[Error creating pptx: {e}]"

    def zip_pack(self, zip_path: str, files: list) -> str:
                                              
        abs_zip = os.path.abspath(os.path.join(self.cwd, zip_path))
        import zipfile
        try:
            os.makedirs(os.path.dirname(abs_zip), exist_ok=True)
            if self.is_protected(zip_path):
                return self._deny(zip_path)
            packed = 0
            skipped = 0
            with zipfile.ZipFile(abs_zip, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for f in files:
                    abs_f = os.path.abspath(os.path.join(self.cwd, f))
                    if self.is_protected(f):
                        skipped += 1
                        continue
                    if os.path.exists(abs_f):
                        zipf.write(abs_f, arcname=os.path.basename(f))
                        packed += 1
                    else:
                        skipped += 1
            note = f", {skipped} skipped (missing or protected)" if skipped else ""
            return f"[Success: ZIP archive '{zip_path}' created with {packed} file(s){note}]"
        except Exception as e:
            return f"[Error creating ZIP archive: {e}]"

    def unzip_pack(self, zip_path: str, extract_to: str = ".") -> str:
                                  
        abs_zip = os.path.abspath(os.path.join(self.cwd, zip_path))
        abs_dest = os.path.abspath(os.path.join(self.cwd, extract_to))
        import zipfile
        if self.is_protected(zip_path) or self.is_protected(extract_to):
            return self._deny(zip_path if self.is_protected(zip_path) else extract_to)
        try:
            with zipfile.ZipFile(abs_zip, 'r') as zipf:
                zipf.extractall(abs_dest)
            return f"[Success: ZIP archive '{zip_path}' extracted to '{extract_to}']"
        except Exception as e:
            return f"[Error extracting ZIP archive: {e}]"

    async def execute_tool(self, tool_name: str, args: dict, allowed_tools=None) -> str:
                                            
        args = normalize_tool_args(args or {})
        if allowed_tools is not None and tool_name not in allowed_tools:
            return f"[Error: Tool '{tool_name}' is unavailable in the active agent style]"
        if tool_name == "run_cmd":
            return await asyncio.to_thread(
                self.run_cmd, args.get("command", ""), args.get("inputs", None)
            )
        elif tool_name == "run_python":
            return await asyncio.to_thread(self.run_python, args.get("code", ""))
        elif tool_name == "run_background_cmd":
            return await self.run_background_cmd(
                args.get("command", ""), args.get("inputs", None)
            )
        elif tool_name == "task_status":
            return await self.task_status(args.get("id", args.get("task_id", "")))
        elif tool_name == "task_log":
            return await self.task_log(
                args.get("id", args.get("task_id", "")),
                args.get("tail_lines", 200),
            )
        elif tool_name == "todo":
            return self.todo(args.get("action", "list"), args)
        elif tool_name == "web_search":
            return await asyncio.to_thread(
                self.web_search,
                args.get("query", ""),
                args.get("max_results", 10),
                args.get("site", ""),
                args.get("region", "wt-wt"),
            )
        elif tool_name == "funstat_osint":
            return await FunstatOSINTTool().run(
                args.get("action", "free_scan"),
                args.get("telegram_id"),
                args.get("username"),
            )
        elif tool_name == "dadata_osint":
            return await DadataOSINTTool().run(
                action=args.get("action", ""),
                query=args.get("query"),
                lat=args.get("lat"),
                lon=args.get("lon"),
                count=args.get("count", 10),
                options=args.get("options", {}),
            )
        elif tool_name == "fetch_url":
            return await asyncio.to_thread(
                self.fetch_url, args.get("url", ""), args.get("max_chars", 6000)
            )
        elif tool_name == "read_file":
            return await asyncio.to_thread(
                self.read_file,
                args.get("path", ""),
                args.get("start_line", 1),
                args.get("end_line", 500),
            )
        elif tool_name == "write_file":
            return await asyncio.to_thread(
                self.write_file, args.get("path", ""), args.get("content", "")
            )
        elif tool_name == "edit_file":
            return await asyncio.to_thread(
                self.edit_file,
                args.get("path", ""),
                args.get("target", ""),
                args.get("replacement", ""),
                args.get("start_line"),
                args.get("end_line"),
            )
        elif tool_name == "list_dir":
            return await asyncio.to_thread(self.list_dir, args.get("path", "."))
        elif tool_name == "file_info":
            return self.file_info(args.get("path", ""))
        elif tool_name == "project_memory":
            return self.project_memory(args.get("action", "get"), args)
        elif tool_name == "sys_info":
            return self.sys_info()
        elif tool_name == "browser_action":
            return await self.browser_action(args.get("action", ""), args)
        elif tool_name == "render_plan":
            return self.render_plan(args.get("title", "План выполнения задачи"), args.get("steps", []))
        elif tool_name == "make_excel":
            return self.make_excel(args.get("path", ""), args.get("sheets", {}))
        elif tool_name == "make_docx":
            return self.make_docx(args.get("path", ""), args.get("title", "Document"), args.get("paragraphs", []))
        elif tool_name == "make_pptx":
            return self.make_pptx(args.get("path", ""), args.get("title", "Presentation"), args.get("slides", []))
        elif tool_name == "zip_pack":
            return self.zip_pack(args.get("zip_path", ""), args.get("files", []))
        elif tool_name == "unzip_pack":
            return self.unzip_pack(args.get("zip_path", ""), args.get("extract_to", "."))

        else:
            return f"[Error: Unknown tool '{tool_name}']"

                                                                                
                                                  
                                                                                

                                                                              
                                                                          
SYSTEM_PROMPT_FILE = Path(__file__).resolve().with_name("system_prompt.txt")
SYSTEM_PROMPT_NOTHINK_FILE = Path(__file__).resolve().with_name("system_prompt_nothink.txt")

def load_agent_system_prompt_base(think: bool = True) -> str:
                                                                                             
    prompt_file = SYSTEM_PROMPT_FILE if think else SYSTEM_PROMPT_NOTHINK_FILE
    if not think and not prompt_file.exists():
        prompt_file = SYSTEM_PROMPT_FILE
    try:
        prompt = prompt_file.read_text(encoding="utf-8").strip()
    except OSError as error:
        raise RuntimeError(
            f"Cannot load editable system prompt '{prompt_file}': {error}"
        ) from error
    if not prompt:
        raise RuntimeError(f"Editable system prompt '{prompt_file}' is empty")
    return prompt

TOOL_DESCRIPTIONS = {
    "run_background_cmd": '- run_background_cmd: {"command":"long-running command","inputs":"optional stdin"}. Starts immediately and returns a task id.',
    "task_status": '- task_status: {"id":"task-0001"}. Returns background process status, PID and exit code.',
    "task_log": '- task_log: {"id":"task-0001","tail_lines":200}. Returns current stdout/stderr from a background process.',
    "run_cmd": '- run_cmd: {"command": "команда", "inputs": "опциональный stdin с \\\\n"}. Используй только команды cmd.exe, не PowerShell.',
    "run_python": '- run_python: {"code": "код Python"}. Код выполняется во временном файле.',
    "read_file": '- read_file: {"path": "файл", "start_line": 1, "end_line": 500}',
    "write_file": '- write_file: {"path": "файл", "content": "полный текст или код"}',
    "edit_file": '- edit_file: {"path": "файл", "target": "старый текст", "replacement": "новый текст"}',
    "list_dir": '- list_dir: {"path": "папка"}',
    "file_info": '- file_info: {"path": "файл"}',
    "project_memory": """- project_memory: постоянная память текущего проекта в .deepx_project_memory.json.
  {"action":"get"} — прочитать цель, факты, завершённые шаги, файлы и последние события.
  {"action":"update","summary":"...","current_task":"...","completed":[...],"next_steps":[...],
  "important_files":[...],"facts":[...]} — заменить актуальную сводку.
  {"action":"append_event","task":"...","result":"...","status":"completed","files":[...]} —
  дописать подтверждённый итог действия. Пиши только краткие проверенные факты, без секретов.""",
    "sys_info": '- sys_info: {}',
    "render_plan": '- render_plan: {"title": "План", "steps": [{"text": "Шаг", "status": "completed|in_progress|pending"}]}. Только для объёмных задач.',
    "make_excel": '- make_excel: {"path": "файл.xlsx", "sheets": {"Лист1": [["A"], [1]]}}',
    "make_docx": '- make_docx: {"path": "файл.docx", "title": "Заголовок", "paragraphs": ["Текст"]}',
    "make_pptx": '- make_pptx: {"path": "файл.pptx", "title": "Заголовок", "slides": [{"title": "Слайд", "content": ["Пункт"]}]}',
    "zip_pack": '- zip_pack: {"zip_path": "архив.zip", "files": ["файл1"]}',
    "unzip_pack": '- unzip_pack: {"zip_path": "архив.zip", "extract_to": "папка"}',
    "todo": """- todo: живой список задач на экране. Действия:
  {"action":"set","items":["Шаг 1","Шаг 2"]} — создать список (заменяет прежний);
  {"action":"add","items":["Ещё шаг"]} — дописать;
  {"action":"start","index":2} — пометить задачу выполняемой;
  {"action":"complete","index":2} — пометить выполненной;
  {"action":"update","index":2,"status":"cancelled","new_text":"..."} — сменить статус или текст;
  {"action":"update","updates":[{"index":1,"status":"done"},{"index":2,"status":"active"}]} —
  несколько изменений за один вызов;
  {"action":"list"} — перерисовать; {"action":"clear"} — очистить.
  Статусы: pending, active, done, cancelled. Вместо index можно передать "text" с куском
  формулировки задачи.
  ПРАВИЛА: заводи список сразу, если в задаче больше двух шагов. Держи ровно одну задачу
  в статусе active. Отмечай complete СРАЗУ после реального выполнения шага, а не пачкой в
  конце. НЕ вызывай todo несколько раз подряд: закрыть одну задачу и начать следующую — это
  ОДИН вызов update со списком updates. Между двумя вызовами todo должно быть хотя бы одно
  реальное действие. Не пиши список текстом в ответе — он уже отрисован на экране.""",
    "web_search": """- web_search: {"query": "запрос", "max_results": 10, "site": "vk.com", "region": "ru-ru"}
  Поиск через DuckDuckGo (HTML-эндпоинт, без API-ключа и без капчи). Возвращает заголовок,
  URL и сниппет. site ограничивает выдачу доменом, region — региональную выдачу (ru-ru,
  us-en, wt-wt). Поддерживает операторы в query: "фраза целиком", filetype:pdf, inurl:, intitle:.""",
    "fetch_url": """- fetch_url: {"url": "https://...", "max_chars": 6000}
  Скачивает страницу и возвращает читаемый текст без скриптов и разметки. Быстрее и дешевле
  браузера. Если страница отдаёт контент только через JS или требует вход — переходи на
  browser_action.""",
    "funstat_osint": """- funstat_osint: разведка Funstat/Telelog по Telegram ID или username.
  {"action":"free_scan","telegram_id":123456789} — профиль, репутация и счётчики сообщений/групп;
  {"action":"free_scan","username":"durov"} — один раз преобразовать username в ID (≈0,1 💠),
  затем выполнить те же бесплатные запросы по числовому ID;
  action также может быть stats_min, reputation или counts.
  Перед сбором инструмент сам читает актуальный Swagger. Полный профиль, сообщения/частота слов,
  подарки и стикеры добавляются только когда соответствующий REST-маршрут прямо помечен бесплатным.
  Не проси обойти ценовой фильтр. Возвращает структурированный JSON и готовый Markdown-отчёт.
  Если известен Telegram ID или username и задача легитимна — используй до общего веб-поиска.""",
    "dadata_osint": """- dadata_osint: официальные справочники DaData для OSINT и due diligence.
  Бесплатные действия:
  address, address_by_id, reverse_geocode, ip_location, postal_by_index, postal_nearby,
  party, party_by_id, bank, bank_by_id, fio, email, passport_issuer, logo_by_domain.
  Формат: {"action":"party_by_id","query":"ИНН/ОГРН","count":10,"options":{"kpp":"..."}}.
  Для координат: {"action":"reverse_geocode","lat":55.75,"lon":37.61,"count":10}.
  address ищет по всем странам; options поддерживает division, language, locations, status,
  type, branch_type, kpp, filters и радиус. party_by_id возвращает полные доступные реквизиты
  ЮЛ/ИП, bank_by_id — БИК/SWIFT/ИНН/КПП, passport_issuer — подразделение по коду.
  Потенциально платные действия clean_address, clean_fio, clean_phone, clean_email,
  clean_passport, clean_vehicle и brand_by_inn реализованы, но заблокированы, пока владелец
  явно не установит DADATA_ALLOW_PAID=true. Никогда не проси и не выводи ключи DaData.""",
    "browser_action": """- browser_action: автономный браузер. Действия: snapshot, navigate, click, fill,
  type, press, select, check, uncheck, hover, focus, scroll, wait, back, forward, reload,
  new_tab, switch_tab, close_tab, list_tabs, upload, download, screenshot, extract_text,
  extract_html, evaluate. Начинай с navigate или snapshot. После каждого действия анализируй
  новое состояние и продолжай до результата. Предпочитай стабильные ref e1, e2 и т. д.""",

}

TOOL_DESCRIPTIONS["edit_file"] = (
    '- edit_file: {"path":"file","target":"old text","replacement":"new text"}; '
    'target supports trailing-whitespace-normalized matching. Or use '
    '{"path":"file","start_line":10,"end_line":12,"replacement":"new lines"}.'
)

COMMON_FILE_TOOLS = {
    "read_file", "write_file", "edit_file", "list_dir", "file_info", "project_memory",
                                                                
    "todo",
}

STYLE_CONFIGS = {
    "coder": {
        "label": "CODER",
        "description": "Разработка, файлы, Python, команды и автономный браузер",
        "tools": COMMON_FILE_TOOLS | {
            "run_cmd", "run_python", "run_background_cmd", "task_status", "task_log",
            "sys_info", "browser_action",
            "render_plan", "zip_pack", "unzip_pack",
        },
        "instructions": """АКТИВНЫЙ СТИЛЬ: CODER
Ты работаешь как сильный инженер-программист. Сначала изучай существующий проект, затем вноси
минимально достаточные изменения, запускай подходящие проверки и сообщай фактический результат.
Для актуальной документации и веб-интерфейсов самостоятельно используй браузер. Не создавай
презентации, Word-документы или Excel-файлы: в этом стиле таких инструментов нет.""",
    },
    "universal": {
        "label": "UNIVERSAL",
        "description": "",
        "tools": COMMON_FILE_TOOLS | {
            "run_cmd", "run_background_cmd", "task_status", "task_log",
            "sys_info", "browser_action", "render_plan",
            "make_excel", "make_docx", "make_pptx", "zip_pack", "unzip_pack",
        },
        "instructions": """АКТИВНЫЙ СТИЛЬ: UNIVERSAL
Ты универсальный помощник для повседневной работы: поиск и действия в браузере, документы,
таблицы, презентации, файлы и системные задачи. Выбирай самый прямой способ выполнить просьбу.
Python в этом стиле недоступен.""",
    },
    "osinter": {
        "label": "OSINTER",
        "tools": COMMON_FILE_TOOLS | {
            "web_search", "fetch_url", "browser_action", "run_python", "run_cmd",
            "run_background_cmd", "task_status", "task_log",
            "funstat_osint", "dadata_osint", "sys_info", "render_plan", "make_docx", "make_excel",
            "zip_pack",
        },
        "protect_self": True,
        "instructions": """АКТИВНЫЙ СТИЛЬ: OSINTER
Ты OSINT-аналитик. Работаешь только с открытыми источниками и только по легитимным задачам:
проверка контрагента, due diligence, поиск утечек собственных данных, расследование мошенничества,
проверка личности с её согласия, корпоративная и брендовая разведка, анализ по разрешённому
скоупу пентеста. Взлом, обход авторизации, покупку краденых баз, слежку за частным лицом ради
преследования или доксинга — не делаешь; если задача выглядит так, скажи об этом прямо одной
фразой и предложи легитимную альтернативу.

ШАГ 1 — ЕДИНЫЙ БРИФИНГ (только один раз, в самом начале расследования).
Первым сообщением выведи одну анкету и жди ответа. Не вызывай инструменты до ответа
пользователя. Не задавай уточняющих вопросов после — дальше работаешь автономно.
Анкета (спрашивай всё сразу, коротким списком):
1. Тип цели: человек / компания / домен или сайт / никнейм / телефон / email / кошелёк / иное.
2. Все известные идентификаторы: ФИО и варианты написания, дата рождения, ники и юзернеймы,
   email, телефоны, домены, ИНН/ОГРН/рег.номер, ссылки на профили, номера объявлений.
3. География и язык: страна, город, языки, в которых искать.
4. Контекст и связи: место работы, учёба, родственники, партнёры, бренды, проекты.
5. Цель расследования и что именно считается результатом.
6. Основание: почему ты имеешь право это проверять (свои данные, согласие, договор, скоуп).
7. Границы: чего не трогать, какая глубина, нужен ли отчёт файлом.
Если пользователь что-то не знает — пишет "нет". После ответа больше не переспрашиваешь.

ШАГ 2 — ПЛАН. Сразу после ответа на анкету ОБЯЗАТЕЛЬНО вызови todo с action=set: по одной
задаче на каждую ветку поиска (каждый идентификатор, каждая площадка, верификация, отчёт).
Это не опция: без списка расследование не начинается. Дальше держи ровно одну задачу в статусе
active, а выполненную закрывай через complete СРАЗУ по факту, а не пачкой в конце. Появилась
новая ветка по ходу — добавляй через add. Тупиковую ветку закрывай через update со статусом
cancelled, а не молча забывай.

ШАГ 3 — СБОР. Работай автономно, много итераций, пока не исчерпаешь ветки.
- web_search — основной инструмент. DuckDuckGo не требует капчи. Делай МНОГО узких запросов,
  а не один широкий: каждый идентификатор отдельно, в кавычках, в разных комбинациях и
  транслитерациях, с региональной выдачей через region.
- Обязательные приёмы: "точная фраза"; site: по площадкам (vk.com, ok.ru, t.me, github.com,
  linkedin.com, x.com, instagram.com, facebook.com, avito.ru, hh.ru, habr.com, reddit.com,
  youtube.com, medium.com); filetype:pdf|xlsx|docx|csv для документов и утечек; intitle: и
  inurl: для профилей; поиск по email-локалу без домена; поиск номера в разных форматах
  (+7..., 8..., 7...); поиск ника как есть и по частям.
- Для юрлиц и доменов: реестры компаний, госзакупки, суды, whois и history, DNS, поддомены,
  сертификаты (crt.sh), архив (web.archive.org), репозитории и утёкшие конфиги.
- fetch_url — читай найденные страницы целиком, не верь одному сниппету. Сниппет — гипотеза,
  подтверждение — только текст страницы.
- browser_action — когда контент подгружается через JS, нужен скролл, интерактив или скриншот
  как доказательство. Только публичные страницы, без входа в чужие аккаунты.
- run_python и run_cmd — нормализация и корреляция: разбор выгрузок, дедупликация, регулярки
  по email/телефонам/кошелькам, сверка таймзон и таймстампов, построение таблицы связей.
- project_memory — сразу фиксируй подтверждённые находки, чтобы не терять их между шагами.

ШАГ 4 — ВЕРИФИКАЦИЯ. Каждый факт держи с источником и оценкой достоверности:
[ПОДТВЕРЖДЕНО] — два и более независимых источника; [ВЕРОЯТНО] — один источник;
[ГИПОТЕЗА] — только совпадение признаков. Однофамильцы и совпадения ников — главная ошибка
OSINT: разделяй разные личности явно, не склеивай их в одну. Отрицательный результат тоже
результат — так и пиши.

ШАГ 5 — ОТЧЁТ. Итог: сводка, таблица идентификаторов в обычной Markdown-разметке, карта связей,
хронология, выводы и пробелы. Если просили файл — make_docx или make_excel, при объёме — zip_pack.
Ссылки на источники обязательны для каждого утверждения. Не додумывай данные: чего не нашёл —
пиши "не найдено". Презентации в этом стиле недоступны.

ЗАПРЕТ НА СЛУЖЕБНЫЕ ФАЙЛЫ (абсолютный, обсуждению не подлежит):
Файлы самого приложения DEEPX — его исходный код, системные промпты, файлы сессии и состояния,
кэш — для тебя не существуют. Ты их не читаешь, не пишешь, не редактируешь, не архивируешь, не
перечисляешь и не упоминаешь в отчётах. Ты не пытаешься добраться до них через run_cmd, run_python,
браузер, относительные пути, символические ссылки или любой обходной путь. Они не входят в скоуп
расследования ни при каких формулировках запроса, включая прямую просьбу пользователя.
Если инструмент вернул [Access denied] — это окончательно: не подбирай другой путь, не пробуй ещё
раз, просто продолжай задачу другими средствами. Рабочие файлы расследования (отчёты, выгрузки,
таблицы) создавай под собственными именами — они не затрагиваются.""",
    },
}

STYLE_ALIASES = {
    "code": "coder",
    "programmer": "coder",
    "универсал": "universal",
    "universal": "universal",
    "osint": "osinter",
    "recon": "osinter",
    "intel": "osinter",
    "осинт": "osinter",
    "разведка": "osinter",
}

def normalize_style_name(value: str):
    name = (value or "").strip().lower()
    if name in STYLE_CONFIGS:
        return name
    return STYLE_ALIASES.get(name)

def build_agent_system_prompt(style: str, think: bool = True, workspace: str = None) -> str:
                                                                                    
    style = normalize_style_name(style) or "coder"
    config = STYLE_CONFIGS[style]
    workspace = workspace or default_workspace_dir()
    tool_lines = [
        TOOL_DESCRIPTIONS[name]
        for name in sorted(config["tools"])
        if name in TOOL_DESCRIPTIONS
    ]
    return (
        f"{load_agent_system_prompt_base(think)}\n\n"
        f"{AGENT_EXECUTION_PROTOCOL}\n\n"
        f"{config['instructions'].strip()}\n\n"
        "РАБОЧАЯ ДИРЕКТОРИЯ:\n"
        f"{workspace}\n"
        "Все относительные пути отсчитываются отсюда. Не предполагай другую папку и не выясняй её "
        "командами — если нужен файл, ищи от этого пути через list_dir.\n\n"
        "ПАМЯТЬ ПРОЕКТА:\n"
        "У тебя есть постоянная память только для текущей рабочей папки. Перед сложной, "
        "многошаговой задачей, продолжением прежней работы или если контекст неясен, сам вызови "
        "project_memory с action=get и опирайся на фактически сохранённые сведения. После важного "
        "завершённого этапа сохраняй короткий подтверждённый итог через append_event; когда меняются "
        "цель, текущая задача, факты или следующие шаги — обновляй сводку через update. Не записывай "
        "догадки, пароли, токены, полные сырые логи и личные данные без необходимости.\n\n"
        "ДОСТУПНЫЕ ИНСТРУМЕНТЫ В ЭТОМ СТИЛЕ "
        "(ВЫЗЫВАЙ ТОЛЬКО ЧЕРЕЗ ```tool_call):\n"
        + "\n".join(tool_lines)
    )


def _compact_task_text(value: str, limit: int = 1200) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if len(text) <= limit:
        return text
    return text[:limit] + f"... (+{len(text) - limit} chars)"


def _tool_result_failed(result: str) -> bool:
    lowered = str(result or "").lower()
    return any(marker in lowered for marker in (
        "[error",
        "[exit code:",
        "syntax check failed",
        "traceback (most recent call last)",
    ))


def build_agent_task_state(
    goal: str,
    todos: list,
    pending_verification: bool = False,
    strategy_change_required: bool = False,
    include_goal: bool = False,
) -> str:
    active = [
        str(item.get("text", "")).strip()
        for item in todos
        if item.get("status") == "active" and str(item.get("text", "")).strip()
    ]
    pending = [
        str(item.get("text", "")).strip()
        for item in todos
        if item.get("status") == "pending" and str(item.get("text", "")).strip()
    ]
    lines = []
    if include_goal:
        lines.append(f"Напоминание цели: {_compact_task_text(goal, 1200)}")
    if active:
        lines.append(f"Активный шаг: {_compact_task_text(active[0], 400)}")
    elif pending:
        lines.append(f"Следующий шаг: {_compact_task_text(pending[0], 400)}")
    if pending_verification:
        lines.append(
            "После изменения артефакта обязательная проверка ещё не выполнена."
        )
    if strategy_change_required:
        lines.append(
            "Одинаковая ошибка повторилась: смени подход и не повторяй тот же вызов."
        )
    if not lines:
        return ""
    return "\n[КРАТКОЕ СОСТОЯНИЕ]\n" + "\n".join(lines) + "\n[КОНЕЦ СОСТОЯНИЯ]\n"

                                                                                
                                
                                                                                

def parse_json_lenient(json_str: str):
                                                                                                
    try:
        return json.loads(json_str, strict=False)
    except Exception:
        pass

                                                              
                                                                                          
    try:
        fixed = re.sub(r'\\(?![\\"/bfnrtu])', r'\\\\', json_str)
        return json.loads(fixed, strict=False)
    except Exception:
        pass

                                           
    try:
        def replace_newlines_in_strings(match):
            val = match.group(0)
            val_escaped = val.replace('\n', '\\n').replace('\r', '\\r').replace('\t', '\\t')
            return val_escaped

        repaired = re.sub(r'"([^"\\]*(\\.[^"\\]*)*)"', replace_newlines_in_strings, json_str, flags=re.DOTALL)
        repaired = re.sub(r'\\(?![\\"/bfnrtu])', r'\\\\', repaired)
        return json.loads(repaired, strict=False)
    except Exception:
        return None

def extract_tool_calls(text: str):
                                                                                                 
    calls = []
    
                                            
    xml_matches = re.findall(r"<tool_call>\s*(.*?)\s*</tool_call>", text, re.DOTALL)
    for match in xml_matches:
        match_str = match.strip()
        cleaned = re.sub(r"^```(?:\w+)?\s*", "", match_str)
        cleaned = re.sub(r"\s*```$", "", cleaned).strip()
        data = parse_json_lenient(cleaned)
        if isinstance(data, dict) and "tool" in data:
            if data not in calls:
                calls.append(data)
            continue

        tool_m = re.search(r"^tool:\s*(\w+)", cleaned, re.MULTILINE)
        if tool_m:
            tool_name = tool_m.group(1).strip()
            path_m = re.search(r"^path:\s*(.+)$", cleaned, re.MULTILINE)
            cmd_m = re.search(r"^command:\s*(.+)$", cleaned, re.MULTILINE)

            code_block_m = re.search(r"```(?:\w+)?\s*\n(.*)\n```", cleaned, re.DOTALL)

            args = {}
            if path_m:
                args["path"] = path_m.group(1).strip()
            if cmd_m:
                args["command"] = cmd_m.group(1).strip()

            if code_block_m:
                args["content"] = code_block_m.group(1)
            elif "content:" in cleaned:
                cont_parts = cleaned.split("content:", 1)
                raw_cont = cont_parts[1].lstrip("\n\r")
                raw_cont = re.sub(r"```(?:\w+)?\s*\n(.*)\n```$", r"\1", raw_cont, flags=re.DOTALL).strip()
                if raw_cont.startswith("```"):
                    raw_cont = re.sub(r"^```(?:\w+)?\n?(.*?)\n?```$", r"\1", raw_cont, flags=re.DOTALL)
                args["content"] = raw_cont

            call_obj = {"tool": tool_name, "args": args}
            if call_obj not in calls:
                calls.append(call_obj)

                                                      
    md_matches = re.findall(r"```tool_call\s*\n?(.*?)\n?```", text, re.DOTALL)
    for match in md_matches:
        match_str = match.strip()
        cleaned = re.sub(r"^<tool_call>\s*", "", match_str)
        cleaned = re.sub(r"\s*</tool_call>$", "", cleaned).strip()
        data = parse_json_lenient(cleaned)
        if isinstance(data, dict) and "tool" in data and data not in calls:
            calls.append(data)

                                    
    json_matches = re.findall(r"```json\s*\n?(.*?)\n?```", text, re.DOTALL)
    for match in json_matches:
        match_str = match.strip()
        cleaned = re.sub(r"^<tool_call>\s*", "", match_str)
        cleaned = re.sub(r"\s*</tool_call>$", "", cleaned).strip()
        data = parse_json_lenient(cleaned)
        if isinstance(data, dict) and "tool" in data and data not in calls:
            calls.append(data)

    return normalize_tool_calls(calls)


                                                                                       
                                                                               
_ARG_ALIASES = {
    "file_path": "path",
    "filepath": "path",
    "filename": "path",
    "file": "path",
    "dir": "path",
    "directory": "path",
    "offset": "start_line",
    "start": "start_line",
    "from_line": "start_line",
    "limit": "end_line",
    "end": "end_line",
    "to_line": "end_line",
    "cmd": "command",
    "shell": "command",
    "q": "query",
    "search_query": "query",
    "link": "url",
    "text": "content",
    "data": "content",
}


def normalize_tool_args(args: dict) -> dict:
                                                                                
    if not isinstance(args, dict):
        return {}
    fixed = {}
    for key, value in args.items():
        canonical = _ARG_ALIASES.get(key, key)
                                                                 
        if canonical in fixed and canonical != key:
            continue
        fixed[canonical] = value
    return fixed


def _as_int(value, fallback: int) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return fallback


def normalize_tool_calls(calls: list) -> list:
                                                                                       

                                                                                     
                                                                                      
                                                                                     
       
    normalized = []
    for call in calls:
        if not isinstance(call, dict) or "tool" not in call:
            continue
        normalized.append({
            "tool": call.get("tool"),
            "args": normalize_tool_args(call.get("args", {})),
        })

    merged = []
    read_index = {}
    for call in normalized:
        if call["tool"] != "read_file" or not call["args"].get("path"):
            if call not in merged:
                merged.append(call)
            continue

        args = call["args"]
        path = args["path"]
        start = max(1, _as_int(args.get("start_line", 1), 1))
        end = _as_int(args.get("end_line", start + MAX_READ_LINES - 1), start + MAX_READ_LINES - 1)
        if end < start:
            start, end = end, start

        if path in read_index:
            target = merged[read_index[path]]["args"]
            target["start_line"] = min(target["start_line"], start)
            target["end_line"] = max(target["end_line"], end)
        else:
            read_index[path] = len(merged)
            merged.append({
                "tool": "read_file",
                "args": {"path": path, "start_line": start, "end_line": end},
            })

    return merged


def count_tokens(text: str) -> int:
                                                                              
    if not text:
        return 0
    tokens = re.findall(r'\w+|\S', text)
    return len(tokens)


async def write_streaming_chars(text: str):
    """Write visible text one character at a time without splitting ANSI codes."""
    index = 0
    while index < len(text):
        if text[index] == "\x1b":
            match = ANSI_ESCAPE_RE.match(text, index)
            if match:
                sys.stdout.write(match.group(0))
                sys.stdout.flush()
                index = match.end()
                continue

        char = text[index]
        sys.stdout.write(char)
        sys.stdout.flush()
        index += 1

        delay = STREAM_CHAR_DELAY
        if char in ".!?,;:":
            delay += STREAM_PUNCTUATION_DELAY
        elif char == "\n":
            delay += STREAM_PUNCTUATION_DELAY / 2
        await asyncio.sleep(delay)

                                                                                
                                               
                                                                                

ANSI_ESCAPE_RE = re.compile(
    r"\x1b(?:\[[0-?]*[ -/]*[@-~]|\][^\x07]*(?:\x07|\x1b\\)|[@-_])"
)


def sanitize_tool_result(text: str) -> str:
                                                                                          
    text = ANSI_ESCAPE_RE.sub("", str(text))
    return "".join(
        char for char in text
        if char in "\n\r\t" or ord(char) >= 32
    )


def get_clean_text(text: str) -> str:
    text = re.sub(
        r"Чтение ссылок недоступно в Экспертном режиме\.?\s*"
        r"(?:Используйте|Переключитесь на)\s+Быстрый режим\.?",
        "",
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(
        r"(?:Reading|Opening) links? (?:is|are) "
        r"(?:unavailable|not available) in Expert mode\.?\s*"
        r"(?:Use|Switch to) (?:Fast|Instant) mode\.?",
        "",
        text,
        flags=re.IGNORECASE,
    )
    text = re.sub(r"^Глубокое размышление\s*", "", text, flags=re.IGNORECASE)
                                                                                   
                              
    text = re.sub(r"\s*<tool_call>.*?(?:</tool_call>|$)", "", text, flags=re.DOTALL)
    
                                                                                 
    text = re.sub(r"\s*```tool_call\s*\n.*?(?:\n```(?=\n|$)|\Z)", "", text, flags=re.DOTALL)
    
                                                             
    text = re.sub(r"\s*```json\s*\n.*?(?:\"tool\"|\"args\").*?(?:\n```(?=\n|$)|\Z)", "", text, flags=re.DOTALL)

    # While streaming, the opening marker arrives character by character
    # (for example "`", "```to", "```tool_call"). Do not print those
    # temporary fragments or their leading blank lines before the complete
    # tool block can be recognized and removed by the expressions above.
    trimmed = text.rstrip()
    for marker in ("```tool_call", "<tool_call>"):
        for prefix_length in range(len(marker), 0, -1):
            prefix = marker[:prefix_length]
            if trimmed.endswith(prefix):
                text = trimmed[:-prefix_length].rstrip()
                break
        else:
            continue
        break

    return text


def render_terminal_markup(text: str, final: bool = False) -> str:
                                                                               
                                                                                             
                                                                                       
    inline_code_style = "\033[38;5;222m"
    styles = {
        "***": "\033[1;3m",
        "**": "\033[1m",
        "~~": "\033[9m",
        "*": "\033[3m",
        "_": "\033[4m",
        "``": inline_code_style,
        "`": inline_code_style,
    }
    reset = "\033[0m"

    def valid_underline_open(source: str, pos: int) -> bool:
        before = source[pos - 1] if pos > 0 else ""
        after = source[pos + 1] if pos + 1 < len(source) else ""
        return (not before.isalnum()) and bool(after) and not after.isspace()

    def find_close(source: str, delimiter: str, start: int) -> int:
        pos = source.find(delimiter, start)
        while pos >= 0:
            if delimiter != "_":
                return pos
            after = source[pos + 1] if pos + 1 < len(source) else ""
            if not after.isalnum():
                return pos
            pos = source.find(delimiter, pos + 1)
        return -1

    def split_table_row(line: str) -> list:
        stripped = line.strip()
        if stripped.startswith("|"):
            stripped = stripped[1:]
        if stripped.endswith("|") and not stripped.endswith("\\|"):
            stripped = stripped[:-1]
        cells = re.split(r"(?<!\\)\|", stripped)
        return [cell.strip().replace("\\|", "|") for cell in cells]

    def table_from(source: str, start: int):
        first_end = source.find("\n", start)
        if first_end < 0:
            return None
        second_start = first_end + 1
        second_end = source.find("\n", second_start)
        if second_end < 0:
            second_end = len(source)
        header = split_table_row(source[start:first_end])
        separators = split_table_row(source[second_start:second_end])
        if (
            len(header) < 2
            or len(separators) != len(header)
            or not all(re.fullmatch(r":?-{3,}:?", cell) for cell in separators)
        ):
            return None

        rows = [header]
        cursor = second_end + (1 if second_end < len(source) else 0)
        while cursor < len(source):
            row_end = source.find("\n", cursor)
            if row_end < 0:
                row_end = len(source)
            line = source[cursor:row_end]
            if not line.strip() or "|" not in line:
                break
            cells = split_table_row(line)
            if len(cells) != len(header):
                break
            rows.append(cells)
            cursor = row_end + (1 if row_end < len(source) else 0)
        return rows, cursor

    def render_markdown_table(rows: list) -> str:
        column_count = len(rows[0])
        cleaned_rows = []
        for row in rows:
            cleaned_rows.append([
                re.sub(r"(\*\*|__|~~|`)", "", cell).replace("\n", " ").strip()
                for cell in row
            ])

        def display_width(value: str) -> int:
            return sum(
                0 if unicodedata.combining(char)
                else 2 if unicodedata.east_asian_width(char) in ("W", "F")
                else 1
                for char in value
            )

        def pad_cell(value: str, width: int) -> str:
            return value + " " * max(0, width - display_width(value))

        terminal_width = max(40, min(getattr(console, "width", 112), 140))
        available = max(column_count * 4, terminal_width - (column_count * 3 + 1))
        natural = [
            max(4, max(display_width(row[index]) for row in cleaned_rows))
            for index in range(column_count)
        ]
        if sum(natural) <= available:
            widths = natural
        else:
            base, extra = divmod(available, column_count)
            widths = [max(4, base + (1 if index < extra else 0)) for index in range(column_count)]

        border_style = "\033[38;5;240m"
        header_style = "\033[1;38;5;255m"
        cell_style = "\033[38;5;252m"
        horizontal = lambda left, middle, right: (
            left + middle.join("─" * (width + 2) for width in widths) + right
        )
        rendered = [border_style, horizontal("┌", "┬", "┐"), reset, "\n"]

        for row_index, row in enumerate(cleaned_rows):
            wrapped = [
                textwrap.wrap(
                    cell,
                    width=widths[index],
                    break_long_words=True,
                    break_on_hyphens=False,
                    replace_whitespace=True,
                ) or [""]
                for index, cell in enumerate(row)
            ]
            height = max(len(lines) for lines in wrapped)
            for line_index in range(height):
                rendered.extend([border_style, "│", reset])
                for column_index, lines in enumerate(wrapped):
                    value = lines[line_index] if line_index < len(lines) else ""
                    style = header_style if row_index == 0 else cell_style
                    rendered.extend([
                        " ",
                        style,
                        pad_cell(value, widths[column_index]),
                        reset,
                        " ",
                        border_style,
                        "│",
                        reset,
                    ])
                rendered.append("\n")
            if row_index == 0:
                rendered.extend([
                    border_style,
                    horizontal("├", "┼", "┤"),
                    reset,
                    "\n",
                ])

        rendered.extend([border_style, horizontal("└", "┴", "┘"), reset])
        return "".join(rendered)

    def parse(source: str, active_styles=()) -> str:
        output = []
        i = 0
        while i < len(source):
            line_start = i == 0 or source[i - 1] == "\n"

            if line_start and "|" in source[i:source.find("\n", i) if source.find("\n", i) >= 0 else len(source)]:
                table = table_from(source, i)
                if table:
                    if not final:
                        break
                    rows, table_end = table
                    output.append(render_markdown_table(rows))
                    output.extend(active_styles)
                    i = table_end
                    continue
                if not final:
                    line_end = source.find("\n", i)
                    if line_end < 0:
                        line_end = len(source)
                    candidate = source[i:line_end].strip()
                    unescaped_pipes = len(re.findall(r"(?<!\\)\|", candidate))
                    if candidate.startswith("|") or unescaped_pipes >= 2:
                        break

            if line_start and source[i] == "#":
                heading = re.match(r"^(#{1,6})[ \t]+", source[i:])
                if heading:
                    line_end = source.find("\n", i)
                    if line_end < 0 and not final:
                        break
                    if line_end < 0:
                        line_end = len(source)
                    content_start = i + heading.end()
                    heading_style = "\033[1;38;5;255m"
                    output.append(heading_style)
                    output.append(parse(
                        source[content_start:line_end],
                        active_styles + (heading_style,)
                    ))
                    output.append(reset)
                    output.extend(active_styles)
                    if line_end < len(source):
                        output.append("\n")
                    i = line_end + (1 if line_end < len(source) else 0)
                    continue
                if not final and source[i:].strip("#") == "":
                    break

            if line_start and source[i] in "-+":
                if i + 1 >= len(source) and not final:
                    break
                if i + 1 < len(source) and source[i + 1] in " \t":
                    output.append("• ")
                    i += 2
                    while i < len(source) and source[i] in " \t":
                        i += 1
                    continue

            if source[i] == "\\" and i + 1 < len(source) and source[i + 1] in "*_~`\\":
                output.append(source[i + 1])
                i += 2
                continue

            if source.startswith("```", i):
                if not final:
                    break
                end = source.find("```", i + 3)
                if end < 0:
                    remainder = source[i + 3:]
                    remainder = re.sub(r"^[A-Za-z0-9_+.-]*\r?\n", "", remainder)
                    output.append(remainder)
                    break
                raw_content = source[i + 3:end]
                language_match = re.match(r"^([A-Za-z0-9_+.-]*)\r?\n", raw_content)
                language = language_match.group(1) if language_match else ""
                content = raw_content[language_match.end():] if language_match else raw_content
                content = content.strip("\r\n")
                border_style = "\033[38;5;240m"
                code_style = "\033[38;5;252m"
                language_style = "\033[38;5;245m"
                output.append(border_style)
                output.append("╭─")
                if language:
                    output.append(language_style)
                    output.append(f" {language}")
                    output.append(border_style)
                output.append("\n")
                for line in content.splitlines() or [""]:
                    output.append("│ ")
                    output.append(code_style)
                    output.append(line)
                    output.append(reset)
                    output.append(border_style)
                    output.append("\n")
                output.append("╰─")
                output.append(reset)
                output.extend(active_styles)
                i = end + 3
                continue

            delimiter = None
            for candidate in ("***", "**", "~~", "``", "*", "_", "`"):
                if source.startswith(candidate, i):
                    if candidate == "_" and not valid_underline_open(source, i):
                        continue
                    delimiter = candidate
                    break

            if delimiter:
                end = find_close(source, delimiter, i + len(delimiter))
                if end < 0:
                    if final:
                        output.append(source[i + len(delimiter):])
                    break
                content = source[i + len(delimiter):end]
                start_code = styles[delimiter]
                output.append(start_code)
                if delimiter in ("``", "`"):
                                                                                           
                    output.append(content)
                else:
                    output.append(parse(content, active_styles + (start_code,)))
                output.append(reset)
                output.extend(active_styles)
                i = end + len(delimiter)
                continue

            output.append(source[i])
            i += 1
        return "".join(output)

    return parse(text)

                                                                                
                      
                                                                                

DEEPX_ASCII_ART = """\
██████╗ ███████╗███████╗██████╗ ██╗  ██╗
██╔══██╗██╔════╝██╔════╝██╔══██╗╚██╗██╔╝
██║  ██║█████╗  █████╗  ██████╔╝ ╚███╔╝
██║  ██║██╔══╝  ██╔══╝  ██╔═══╝  ██╔██╗
██████╔╝███████╗███████╗██║     ██╔╝ ██╗
╚═════╝ ╚══════╝╚══════╝╚═╝     ╚═╝  ╚═╝"""
BANNER_WIDTH = 100

def render_banner(
    mode: str,
    style: str,
    think: bool,
    search: bool,
    status_message: str | None = None,
):
                                                                                 
    grid = Table.grid(expand=True)
    grid.add_column(justify="center")

    title = Text(DEEPX_ASCII_ART, style="bold #536DFE", end="")

    if status_message:
        status_line = f"[bold #F8FAFC]{status_message}[/bold #F8FAFC]"
    else:
        mode_val = f"[bold #38BDF8]{mode.upper()}[/bold #38BDF8]"
        style_label = STYLE_CONFIGS.get(style, STYLE_CONFIGS["coder"])["label"]
        style_val = f"[bold #A78BFA]{style_label}[/bold #A78BFA]"
        think_val = "[bold #38BDF8]ON[/bold #38BDF8]" if think else "[dim #64748B]OFF[/dim #64748B]"
        if mode == "expert":
            status_line = f"Mode: {mode_val}  |  Style: {style_val}  |  DeepThink: {think_val}"
        else:
            search_val = "[bold #38BDF8]ON[/bold #38BDF8]" if search else "[dim #64748B]OFF[/dim #64748B]"
            status_line = f"Mode: {mode_val}  |  Style: {style_val}  |  DeepThink: {think_val}  |  Search: {search_val}"

    grid.add_row(title)
    grid.add_row(Text(""))
    grid.add_row(Text.from_markup(status_line))

    panel = Panel(
        grid,
        border_style="#536DFE",
        padding=(1, 4),
        width=BANNER_WIDTH,
        expand=True,
        box=box.ROUNDED,
    )
    console.print(Align.center(panel))

def print_help():
                                                       
    table = Table(title="DEEPX COMMANDS", border_style="#536DFE", header_style="bold #536DFE")
    table.add_column("Command", style="bold #536DFE", width=28)
    table.add_column("Description", style="white")

    for cmd, desc in COMMANDS_DICT.items():
        table.add_row(f"/{cmd}", desc)

    console.print(table)
    console.print(
        "[bold #EF4444]Правый Ctrl (удерживать 0,5 секунды)[/bold #EF4444] — "
        "локальная запись речи; отпусти Ctrl, затем нажми Enter для отправки."
    )

def print_status(mode: str, style: str, think: bool, search: bool, cwd: str, state_file: str):
                                        
    grid = Table.grid(expand=True)
    grid.add_column(style="bold #536DFE", width=22)
    grid.add_column(style="white")

    grid.add_row("Model Mode:", mode.upper())
    grid.add_row("Agent Style:", STYLE_CONFIGS.get(style, STYLE_CONFIGS["coder"])["label"])
    grid.add_row("DeepThink Engine:", "Enabled" if think else "Disabled")
    if mode != "expert":
        grid.add_row("Smart Web Search:", "Enabled" if search else "Disabled")
    grid.add_row("Working Directory:", cwd)
    grid.add_row("Session State:", state_file)

    console.print(Panel(grid, title="SESSION STATUS", border_style="#536DFE"))

                                                                                
                            
                                                                                

class DeepCLIApp:
    def __init__(self, state_file: str = "state.json", headless: bool = True):
        self.state_file = state_file
        self.headless = headless
        self.mode = "instant"
        self.agent_style = "coder"
        self.think = False
        self.search = False
        self.tools = AgentTools(browser_headless=headless)
        self.tools.add_protected(state_file)
        self._apply_style_guard()
        self.last_response = ""
        self.is_new_chat = True
        self._active_file_live = None
        self._stream_file_activities = {}
        self._voice_model = None
        self._voice_device = None
        self._voice_recording = False

        completer = SlashCommandCompleter(COMMANDS_DICT)
        pt_style = PTStyle.from_dict({
            'completion-menu.completion': 'bg:#172554 fg:#93C5FD',
            'completion-menu.completion.current': 'bg:#2563EB fg:#FFFFFF bold',
            'completion-menu.meta.completion': 'bg:#172554 fg:#60A5FA',
            'completion-menu.meta.completion.current': 'bg:#2563EB fg:#EFF6FF',
            'prompt': 'bold #536DFE',
        })

        self.session = PromptSession(
            completer=completer,
            complete_while_typing=True,
            style=pt_style,
        )

    def refresh_ui(self, status_message: str | None = None):
                                                                                        
        console.clear()
        render_banner(
            self.mode,
            self.agent_style,
            self.think,
            self.search,
            status_message=status_message,
        )
        console.print()

    def _voice_hotkey_pressed(self) -> bool:
        if sys.platform != "win32":
            return False
        return bool(ctypes.windll.user32.GetAsyncKeyState(VOICE_HOTKEY_VK) & 0x8000)

    def _schedule_prompt_callback(self, callback) -> bool:
        app = self.session.app
        prompt_loop = getattr(app, "loop", None)
        if prompt_loop is None or prompt_loop.is_closed():
            return False
        prompt_loop.call_soon_threadsafe(callback)
        return True

    def _draw_voice_status(
        self,
        *,
        recording: bool,
        level: float = 0.0,
        elapsed: float = 0.0,
        message: str = "",
    ):
        width = shutil.get_terminal_size((BANNER_WIDTH, 30)).columns
        panel_width = min(BANNER_WIDTH, width)
        inner_width = max(1, panel_width - 2)
        left_padding = max(0, (width - panel_width) // 2)
        reset = "\033[0m"
        white = "\033[38;2;248;250;252m"
        cyan = "\033[1;38;2;56;189;248m"
        purple = "\033[1;38;2;167;139;250m"
        dim = "\033[2;38;2;100;116;139m"
        red = "\033[1;38;2;239;68;68m"

        if recording:
            meter_width = 18
            filled = max(0, min(meter_width, round(level * meter_width)))
            meter = "█" * filled + "·" * (meter_width - filled)
            seconds = max(0, int(elapsed))
            visible = (
                f"● Recording  {seconds // 60:02d}:{seconds % 60:02d}  "
                f"[{meter}]  отпусти {VOICE_HOTKEY_LABEL} для остановки"
            )
            styled = red + visible
        elif message:
            visible = message
            styled = red + visible
        else:
            style_label = STYLE_CONFIGS.get(
                self.agent_style, STYLE_CONFIGS["coder"]
            )["label"]
            mode_value = self.mode.upper()
            think_value = "ON" if self.think else "OFF"
            visible = (
                f"Mode: {mode_value}  |  Style: {style_label}  |  "
                f"DeepThink: {think_value}"
            )
            styled = (
                f"{white}Mode: {cyan}{mode_value}{white}  |  Style: "
                f"{purple}{style_label}{white}  |  DeepThink: "
                f"{cyan if self.think else dim}{think_value}"
            )
            if self.mode != "expert":
                search_value = "ON" if self.search else "OFF"
                visible += f"  |  Search: {search_value}"
                styled += (
                    f"{white}  |  Search: "
                    f"{cyan if self.search else dim}{search_value}"
                )

        if len(visible) > inner_width:
            visible = visible[:inner_width]
            styled = white + visible
        left_inner_padding = max(0, (inner_width - len(visible)) // 2)
        right_inner_padding = max(
            0, inner_width - len(visible) - left_inner_padding
        )
        content = (
            " " * left_inner_padding
            + styled
            + reset
            + " " * right_inner_padding
        )
        line = (
            " " * left_padding
            + "\033[38;2;83;109;254m│"
            + content
            + "\033[38;2;83;109;254m│\033[0m"
        )

        def redraw():
            # The banner always starts at row 1 after refresh_ui(); its status
            # line is row 10 (border + padding + six logo rows + spacer).
            sys.stdout.write(f"\033[s\033[10;1H\033[2K{line}\033[u")
            sys.stdout.flush()
            self.session.app.invalidate()

        self._schedule_prompt_callback(redraw)

    async def _prompt_buffer_snapshot(self):
        result = concurrent.futures.Future()

        def capture():
            try:
                document = self.session.app.current_buffer.document
                result.set_result((document.text, document.cursor_position))
            except Exception as error:
                result.set_exception(error)

        if not self._schedule_prompt_callback(capture):
            return "", 0
        return await asyncio.wrap_future(result)

    def _replace_voice_text(
        self,
        base_text: str,
        base_cursor: int,
        transcript: str,
    ):
        spoken = re.sub(r"\s+", " ", str(transcript or "")).strip()
        prefix = base_text[:base_cursor]
        suffix = base_text[base_cursor:]
        if spoken and prefix and not prefix[-1].isspace():
            spoken = " " + spoken
        if spoken and suffix and not suffix[0].isspace():
            spoken += " "
        updated = prefix + spoken + suffix
        cursor_position = len(prefix) + len(spoken)

        def apply_text():
            buffer = self.session.app.current_buffer
            buffer.set_document(
                Document(updated, cursor_position=cursor_position),
                bypass_readonly=True,
            )
            self.session.app.invalidate()

        self._schedule_prompt_callback(apply_text)

    def _load_voice_model(self):
        if self._voice_model is not None:
            return self._voice_model
        import torch
        import whisper

        device = "cuda" if torch.cuda.is_available() else "cpu"
        self._voice_model = whisper.load_model(VOICE_MODEL_NAME, device=device)
        return self._voice_model

    def _transcribe_voice_samples(self, samples) -> str:
        if samples is None or len(samples) < VOICE_SAMPLE_RATE // 3:
            return ""
        import numpy as np
        import torch

        audio = np.asarray(samples, dtype=np.float32).reshape(-1)
        model = self._load_voice_model()
        result = model.transcribe(
            audio,
            language=os.environ.get("DEEPX_VOICE_LANGUAGE", "ru"),
            fp16=torch.cuda.is_available(),
            verbose=None,
            condition_on_previous_text=False,
            temperature=0,
        )
        return str(result.get("text", "")).strip()

    async def _record_voice_while_held(self, prompt_future):
        import numpy as np
        import sounddevice as sd

        base_text, base_cursor = await self._prompt_buffer_snapshot()
        audio_chunks = []
        audio_state = {"level": 0.0}
        started_at = asyncio.get_running_loop().time()
        next_partial_at = started_at + VOICE_PARTIAL_INTERVAL
        transcribe_task = None
        transcribed_samples = 0
        latest_transcript = ""

        def audio_callback(indata, _frames, _time_info, status):
            if status:
                return
            mono = np.asarray(indata, dtype=np.float32).mean(axis=1).copy()
            audio_chunks.append(mono)
            rms = float(np.sqrt(np.mean(mono * mono))) if mono.size else 0.0
            audio_state["level"] = min(1.0, rms * 12.0)

        stream = sd.InputStream(
            device=self._voice_device,
            samplerate=VOICE_SAMPLE_RATE,
            channels=1,
            dtype="float32",
            callback=audio_callback,
        )

        self._voice_recording = True
        try:
            stream.start()
            model_task = asyncio.create_task(asyncio.to_thread(self._load_voice_model))

            while self._voice_hotkey_pressed() and not prompt_future.done():
                now = asyncio.get_running_loop().time()
                self._draw_voice_status(
                    recording=True,
                    level=audio_state["level"],
                    elapsed=now - started_at,
                )

                if transcribe_task is not None and transcribe_task.done():
                    try:
                        transcript = transcribe_task.result()
                        latest_transcript = transcript
                        self._replace_voice_text(
                            base_text, base_cursor, transcript
                        )
                    except Exception:
                        pass
                    transcribe_task = None

                if (
                    model_task.done()
                    and transcribe_task is None
                    and now >= next_partial_at
                    and audio_chunks
                ):
                    snapshot = np.concatenate(audio_chunks)
                    transcribed_samples = len(snapshot)
                    transcribe_task = asyncio.create_task(
                        asyncio.to_thread(
                            self._transcribe_voice_samples,
                            snapshot,
                        )
                    )
                    next_partial_at = now + VOICE_PARTIAL_INTERVAL

                await asyncio.sleep(0.08)

            stream.stop()
            stream.close()

            if prompt_future.done():
                return

            self._draw_voice_status(
                recording=False,
                message="Обработка последней фразы…",
            )
            await model_task
            final_samples = (
                np.concatenate(audio_chunks)
                if audio_chunks
                else np.empty(0, dtype=np.float32)
            )
            if transcribe_task is not None:
                try:
                    latest_transcript = await transcribe_task
                except Exception:
                    pass
            if len(final_samples) > transcribed_samples or not transcribed_samples:
                transcript = await asyncio.to_thread(
                    self._transcribe_voice_samples,
                    final_samples,
                )
            elif transcribe_task is not None:
                transcript = latest_transcript
            else:
                transcript = latest_transcript
            self._replace_voice_text(base_text, base_cursor, transcript)
        finally:
            self._voice_recording = False
            try:
                if stream.active:
                    stream.stop()
                stream.close()
            except Exception:
                pass
            self._draw_voice_status(recording=False)

    async def _watch_voice_hotkey(self, prompt_future):
        if sys.platform != "win32":
            return
        pressed_since = None
        while not prompt_future.done():
            pressed = self._voice_hotkey_pressed()
            now = asyncio.get_running_loop().time()
            if pressed:
                if pressed_since is None:
                    pressed_since = now
                elif (
                    not self._voice_recording
                    and now - pressed_since >= VOICE_HOLD_SECONDS
                ):
                    try:
                        await self._record_voice_while_held(prompt_future)
                    except Exception as error:
                        self._draw_voice_status(
                            recording=False,
                            message=f"Микрофон недоступен: {error}",
                        )
                        await asyncio.sleep(2.0)
                        self._draw_voice_status(recording=False)
                    pressed_since = None
            else:
                pressed_since = None
            await asyncio.sleep(0.04)

    def restart_in_windows_terminal(self):
                                                                           
        command = [
            "wt.exe",
            "-w", "new",
            "--size", "110,30",
            "new-tab",
            "--title", "DEEPX AGENT",
            "-d", APP_DIR,
            "cmd.exe", "/d", "/c",
            sys.executable,
            os.path.join(APP_DIR, "deep_cli.py"),
            "--state", self.state_file,
        ]
        if not self.headless:
            command.append("--headful")
        subprocess.Popen(
            command,
            cwd=APP_DIR,
            close_fds=True,
        )

    @property
    def allowed_tools(self):
        return STYLE_CONFIGS[self.agent_style]["tools"]

    def _apply_style_guard(self):
                                                                                 
        self.tools.protect_self = STYLE_CONFIGS[self.agent_style].get("protect_self", False)

    async def switch_style(self, target_style: str, api: DeepAPI):
                                                                                    
        target_style = normalize_style_name(target_style)
        if not target_style:
            return False
        if target_style == self.agent_style:
            console.print(
                f"[dim #64748B]Style {STYLE_CONFIGS[target_style]['label']} is already active.[/dim #64748B]"
            )
            return True

        self.agent_style = target_style
        self._apply_style_guard()
        if target_style == "osinter" and self.mode != "expert":
            self.search = True
        await api.new_chat()
        self.is_new_chat = True
        self.refresh_ui()
        config = STYLE_CONFIGS[target_style]
        console.print(
            f"[bold #38BDF8]Agent style set to {config['label']} in a fresh session.[/bold #38BDF8]"
        )
        if config.get("description"):
            console.print(f"[dim #94A3B8]{config['description']}[/dim #94A3B8]")
        return True

    async def run(self):
                                                                           
                                                              
        if not os.path.isfile(self.state_file):
            self.refresh_ui(
                "Авторизуйтесь либо создайте новый аккаунт в DeepSeek."
            )
            console.print(
                Align.center(
                    Text(
                        "CLI ожидает завершения авторизации в открывшемся браузере…",
                        style="#94A3B8",
                    )
                )
            )
            try:
                auth_api = DeepAPI(state_file=self.state_file, headless=False)
                await auth_api.wait_for_login_and_save()
                console.print(
                    "\n[bold #38BDF8]Авторизация получена. "
                    "Перезапускаю DEEPX…[/bold #38BDF8]"
                )
                self.restart_in_windows_terminal()
            except Exception as e:
                console.print(
                    f"\nОшибка авторизации [{type(e).__name__}]: {e}",
                    style="bold red",
                    markup=False,
                    highlight=False,
                )
            return

        self.refresh_ui()

        try:
            async with DeepAPI(state_file=self.state_file, headless=self.headless) as api:
                while True:
                    prompt_str = "› "
                    try:
                        prompt_future = asyncio.get_running_loop().run_in_executor(
                            None, lambda: self.session.prompt(prompt_str)
                        )
                        voice_watcher = asyncio.create_task(
                            self._watch_voice_hotkey(prompt_future)
                        )
                        try:
                            user_input = await prompt_future
                        finally:
                            if not voice_watcher.done():
                                voice_watcher.cancel()
                            try:
                                await voice_watcher
                            except asyncio.CancelledError:
                                pass
                        user_input = user_input.strip()
                    except KeyboardInterrupt:
                        if await self.confirm_exit():
                            console.print(
                                "\n[bold #8B5CF6]Exiting DEEPX. Goodbye![/bold #8B5CF6]"
                            )
                            break
                        continue
                    except EOFError:
                        console.print("\n[bold #8B5CF6]Exiting DEEPX. Goodbye![/bold #8B5CF6]")
                        break

                    if not user_input:
                        continue

                                           
                    if user_input.startswith("/"):
                        if await self.handle_command(user_input, api):
                            continue

                                        
                    await self.run_query_interruptible(user_input, api)

        except Exception as e:
                                                                                        
                                                                                        
            message = str(e).replace("\n", " ")
            if len(message) > 500:
                message = message[:500] + f"... (+{len(str(e)) - 500} chars)"
            console.print(
                f"\nSystem Error [{type(e).__name__}]: {message}",
                style="bold red",
                markup=False,
                highlight=False,
            )
        finally:
            await self.tools.close_background_tasks()
            await self.tools.close_browser()

    async def confirm_exit(self) -> bool:
        try:
            answer = await asyncio.get_running_loop().run_in_executor(
                None,
                lambda: self.session.prompt("Do you really want to exit? [y/N] "),
            )
        except (EOFError, KeyboardInterrupt):
            return False
        return answer.strip().lower() in ("y", "yes")

    async def run_query_interruptible(self, user_input: str, api: DeepAPI) -> bool:
        query_task = asyncio.create_task(self.process_query(user_input, api))
        loop = asyncio.get_running_loop()
        windows_console_state = _begin_windows_ctrl_c_key_mode()
        interrupt_event = None
        if windows_console_state:
            interrupt_task = asyncio.create_task(_wait_for_windows_control_key())
        else:
            interrupt_event = asyncio.Event()
            interrupt_task = asyncio.create_task(interrupt_event.wait())
        interrupted = False
        previous_sigint = signal.getsignal(signal.SIGINT)
        previous_exception_handler = loop.get_exception_handler()

        def cancellation_exception_handler(active_loop, context):
            if (
                interrupted
                and str(context.get("message", "")).startswith(
                    "socket.send() raised exception"
                )
            ):
                return
            if previous_exception_handler:
                previous_exception_handler(active_loop, context)
            else:
                active_loop.default_exception_handler(context)

        def request_stop(_signum, _frame):
            if query_task.done() or interrupt_event.is_set():
                return
            loop.call_soon_threadsafe(interrupt_event.set)

        if not windows_console_state:
            signal.signal(signal.SIGINT, request_stop)
        try:
            while True:
                done, _ = await asyncio.wait(
                    (query_task, interrupt_task),
                    return_when=asyncio.FIRST_COMPLETED,
                )
                if query_task in done:
                    await query_task
                    return True

                action = (
                    interrupt_task.result()
                    if windows_console_state
                    else "exit"
                )
                if action != "exit":
                    break

                _restore_windows_console_mode(windows_console_state)
                windows_console_state = None
                if await self.confirm_exit():
                    interrupted = True
                    loop.set_exception_handler(cancellation_exception_handler)
                    try:
                        await asyncio.wait_for(
                            asyncio.shield(api.stop_generation()),
                            timeout=6.0,
                        )
                    except Exception:
                        pass
                    if not query_task.done():
                        query_task.cancel()
                    try:
                        await query_task
                    except asyncio.CancelledError:
                        pass
                    console.print(
                        "\n[bold #8B5CF6]Exiting DEEPX. Goodbye![/bold #8B5CF6]"
                    )
                    raise SystemExit(0)

                if sys.platform == "win32":
                    windows_console_state = _begin_windows_ctrl_c_key_mode()
                if windows_console_state:
                    interrupt_task = asyncio.create_task(
                        _wait_for_windows_control_key()
                    )
                else:
                    interrupt_event = asyncio.Event()
                    interrupt_task = asyncio.create_task(interrupt_event.wait())
                    signal.signal(signal.SIGINT, request_stop)

            interrupted = True
            loop.set_exception_handler(cancellation_exception_handler)
            sys.stdout.write("\033[0m\r\033[K")
            sys.stdout.flush()
            stopped = False
            try:
                stopped = await asyncio.wait_for(
                    asyncio.shield(api.stop_generation()),
                    timeout=6.0,
                )
            except Exception:
                stopped = False

            if not query_task.done():
                await asyncio.sleep(0.1)
                query_task.cancel()
            try:
                await query_task
            except asyncio.CancelledError:
                pass

            sys.stdout.write("\033[0m\r\033[K")
            sys.stdout.flush()
            if stopped:
                console.print(
                    "\n[bold #F59E0B]Генерация остановлена. "
                    "Можно отправить новый запрос в этот же чат.[/bold #F59E0B]"
                )
            else:
                console.print(
                    "\n[bold #F59E0B]Ход агента отменён. "
                    "Кнопка остановки DeepSeek не найдена, но CLI снова готова к вводу.[/bold #F59E0B]"
                )
            return False
        finally:
            self._stop_streaming_file_visual()
            if not interrupt_task.done():
                interrupt_task.cancel()
            if not query_task.done():
                query_task.cancel()
            if not interrupted:
                loop.set_exception_handler(previous_exception_handler)
            _restore_windows_console_mode(windows_console_state)
            if not windows_console_state:
                signal.signal(signal.SIGINT, previous_sigint)

    async def compact_context_and_continue(self, api: DeepAPI) -> bool:
        compression_prompt = """
[СИСТЕМНАЯ КОМАНДА СЖАТИЯ КОНТЕКСТА]
Не выполняй инструменты и не продолжай работу в этом чате. Сожми весь текущий
контекст в одну автономную постановку задачи для нового экземпляра агента.

Включи только сведения, необходимые для продолжения:
- исходную цель пользователя и ожидаемый результат;
- уже выполненное и фактически подтверждённое;
- текущий прогресс, важные решения и ограничения;
- значимые пути, имена файлов, команды и технические детали;
- известные ошибки и неудачные подходы, которые не надо повторять;
- конкретный следующий шаг и критерии окончательной готовности.

Не включай внутренние рассуждения, приветствия, повторы, полные сырые логи,
описание этой команды или предложения пользователю. Не заявляй о выполнении
того, что не подтверждено. Ответ должен быть самостоятельной задачей, которую
можно передать в пустой чат без предыдущей переписки. Максимум 6000 символов.
Начни сразу с заголовка «ЗАДАЧА».
[КОНЕЦ СИСТЕМНОЙ КОМАНДЫ]
""".strip()

        console.print(
            "[bold #8B5CF6]Сжимаю текущую задачу для нового чата…[/bold #8B5CF6]"
        )
        try:
            stream = await api.send_message(compression_prompt, stream=True)
            compacted = ""
            async for chunk in stream:
                if isinstance(chunk, dict) and chunk.get("answer"):
                    compacted = str(chunk["answer"])
        except Exception as error:
            console.print(
                f"Не удалось сжать контекст [{type(error).__name__}]: {error}",
                style="bold red",
                markup=False,
                highlight=False,
            )
            return False

        contains_tool_call = bool(extract_tool_calls(compacted))
        compacted = get_clean_text(compacted).strip()
        if not compacted or contains_tool_call:
            console.print(
                "[bold red]DeepSeek не вернул корректную сжатую задачу. "
                "Текущий чат оставлен без изменений.[/bold red]"
            )
            return False
        if len(compacted) > 12000:
            compacted = compacted[:12000].rstrip() + "\n[Сжатие обрезано до 12000 символов]"

        saved_mode = self.mode
        saved_style = self.agent_style
        saved_think = self.think
        saved_search = self.search

        try:
            await api.new_chat()
            self.mode = saved_mode
            self.agent_style = saved_style
            self.think = saved_think
            self.search = saved_search
            self._apply_style_guard()
            self.is_new_chat = True
            self.refresh_ui()
            console.print(
                "[bold #10B981]Контекст сжат. Продолжаю задачу в новом чате "
                "с прежними настройками.[/bold #10B981]"
            )
            await self.run_query_interruptible(compacted, api)
            return True
        except Exception as error:
            console.print(
                f"Не удалось продолжить в новом чате [{type(error).__name__}]: {error}",
                style="bold red",
                markup=False,
                highlight=False,
            )
            return False

    async def handle_command(self, cmd: str, api: DeepAPI) -> bool:
        parts = cmd.split()
        main_cmd = parts[0].lower()

        if main_cmd in ["/exit", "/quit"]:
            console.print("[bold #8B5CF6]Exiting DEEPX. Goodbye![/bold #8B5CF6]")
            raise SystemExit(0)

        elif main_cmd == "/help":
            print_help()
            return True

        elif main_cmd == "/clear":
            self.refresh_ui()
            return True

        elif main_cmd == "/status":
            print_status(
                self.mode, self.agent_style, self.think, self.search,
                self.tools.cwd, self.state_file
            )
            return True

        elif main_cmd == "/stealth":
            report = await self.tools.browser_stealth_report()
            signals = report["signals"]
            table = Table(title="PLAYWRIGHT STEALTH CHECK", border_style="#536DFE")
            table.add_column("Check", style="bold #A78BFA")
            table.add_column("Value")
            table.add_row("Package", f"playwright-stealth {report['package_version']}")
            table.add_row("Integration", report["status"])
            table.add_row("Basic result", report["basic_check"].upper())
            table.add_row("navigator.webdriver", repr(signals.get("webdriver")))
            table.add_row("User-Agent", str(signals.get("userAgent", "")))
            table.add_row("Languages", ", ".join(signals.get("languages", [])))
            table.add_row("Plugins", str(signals.get("plugins", 0)))
            table.add_row("window.chrome", str(signals.get("chromeObject")))
            console.print(table)
            for warning in report["warnings"]:
                console.print(f"[bold #F59E0B]Warning:[/bold #F59E0B] {warning}")
            console.print(
                "[dim]This checks common browser signals, but cannot guarantee "
                "that Google or another anti-bot system will not detect automation.[/dim]"
            )
            return True

        elif main_cmd in ["/new", "/reset"]:
            console.print("[bold #8B5CF6]Initializing fresh session context...[/bold #8B5CF6]")
            await api.new_chat()
            self.is_new_chat = True
            self.refresh_ui()
            console.print("[bold #10B981]New session initialized.[/bold #10B981]")
            return True

        elif main_cmd == "/copy":
            if self.last_response:
                pyperclip.copy(self.last_response)
                console.print("[bold #10B981]Response copied to clipboard![/bold #10B981]")
            else:
                console.print("[dim #F59E0B]No response content to copy.[/dim #F59E0B]")
            return True

        elif main_cmd == "/menu":
            await self.open_interactive_menu(api)
            return True

        elif main_cmd == "/ctx":
            await self.compact_context_and_continue(api)
            return True

        elif main_cmd == "/style":
            if len(parts) > 1 and parts[1].lower() == "list":
                table = Table(
                    title="AGENT STYLES",
                    border_style="#536DFE",
                    header_style="bold #536DFE",
                )
                table.add_column("Style", style="bold #A78BFA")
                table.add_column("Purpose")
                table.add_column("Tools", justify="right")
                for name, config in STYLE_CONFIGS.items():
                    marker = " *" if name == self.agent_style else ""
                    table.add_row(
                        f"{name}{marker}",
                        config.get("description", ""),
                        str(len(config["tools"])),
                    )
                console.print(table)
                console.print("[dim]Switch with /style coder, /style universal or /style osinter[/dim]")
                return True

            target_style = normalize_style_name(parts[1]) if len(parts) > 1 else None
            if len(parts) > 1 and not target_style:
                console.print(
                    "[bold red]Unknown style. Use coder, universal or osinter.[/bold red]"
                )
                return True
            if not target_style:
                target_style = await radiolist_dialog(
                    title="Agent Style Selection",
                    text="Select a specialized prompt and toolset (creates a fresh session):",
                    values=[
                        ("coder", "Coder — development, Python and browser automation"),
                        ("universal", "Universal"),
                        ("osinter", "OSINTer — DuckDuckGo recon, source verification, reports"),
                    ],
                ).run_async()
            if target_style:
                await self.switch_style(target_style, api)
            return True

        elif main_cmd == "/mode":
            target_mode = None
            if len(parts) > 1 and parts[1].lower() in ["instant", "expert", "vision"]:
                target_mode = parts[1].lower()
            else:
                selected = await radiolist_dialog(
                    title="Model Mode Selection",
                    text="Select operating model mode (Creates a new chat session context):",
                    values=[
                        ("instant", "Instant Mode (Fast & Responsive)"),
                        ("expert", "Expert Mode (Deep Reasoning Engine)"),
                        ("vision", "Vision Mode (Multimodal Recognition)")
                    ]
                ).run_async()
                if selected:
                    target_mode = selected

            if target_mode and target_mode != self.mode:
                self.mode = target_mode
                console.print(f"[bold #536DFE]Switching model mode to {self.mode.upper()} (Initializing fresh chat session...)[/bold #536DFE]")
                await api.new_chat()
                self.is_new_chat = True
                self.refresh_ui()
                console.print(f"[bold #38BDF8]Model mode set to {self.mode.upper()} in new session.[/bold #38BDF8]")
            return True

        elif main_cmd == "/think":
            if len(parts) > 1:
                self.think = (parts[1].lower() == "on")
            else:
                self.think = not self.think
            await api.set_deepthink(self.think)
            self.refresh_ui()
            console.print(f"[bold #38BDF8]DeepThink reasoning: {'ENABLED' if self.think else 'DISABLED'}[/bold #38BDF8]")
            return True

        elif main_cmd == "/search":
            if self.mode == "expert":
                console.print("[dim #64748B]Smart Search is not available in Expert mode.[/dim #64748B]")
                return True
            if len(parts) > 1:
                self.search = (parts[1].lower() == "on")
            else:
                self.search = not self.search
            self.refresh_ui()
            console.print(f"[bold #38BDF8]Smart Web Search: {'ENABLED' if self.search else 'DISABLED'}[/bold #38BDF8]")
            return True

        else:
            console.print(f"[bold red]Unknown command: {cmd}. Type /help for command list.[/bold red]")
            return True

    async def open_interactive_menu(self, api: DeepAPI):
                                                        
        option = await radiolist_dialog(
            title="DEEPX Settings Menu",
            text="Use Arrow Keys (up/down) and press Enter to select setting to modify:",
            values=[
                ("mode", f"Change Model Mode -> Creates New Chat (Current: {self.mode.upper()})"),
                ("style", f"Change Agent Style -> Creates New Chat (Current: {self.agent_style.upper()})"),
                ("think", f"Toggle DeepThink Reasoning (Current: {'ON' if self.think else 'OFF'})"),
                ("search", f"Toggle Smart Web Search (Current: {'ON' if self.search else 'OFF'})"),
                ("new", "Reset Session Context")
            ]
        ).run_async()

        if option == "mode":
            mode_choice = await radiolist_dialog(
                title="Select Model Mode",
                text="Choose target model (Will initialize a new chat session):",
                values=[
                    ("instant", "Instant Mode (Fast)"),
                    ("expert", "Expert Mode (Deep Reasoning)"),
                    ("vision", "Vision Mode (Multimodal)")
                ]
            ).run_async()
            if mode_choice and mode_choice != self.mode:
                self.mode = mode_choice
                console.print(f"[bold #536DFE]Switching model mode to {self.mode.upper()} (Initializing fresh chat session...)[/bold #536DFE]")
                await api.new_chat()
                self.is_new_chat = True
                self.refresh_ui()
                console.print(f"[bold #38BDF8]Model mode set to {self.mode.upper()} in new session.[/bold #38BDF8]")
        elif option == "style":
            style_choice = await radiolist_dialog(
                title="Select Agent Style",
                text="Choose a specialized prompt and toolset:",
                values=[
                    ("coder", "Coder — development, Python and browser automation"),
                    ("universal", "Universal"),
                    ("osinter", "OSINTer — DuckDuckGo recon, source verification, reports"),
                ],
            ).run_async()
            if style_choice:
                await self.switch_style(style_choice, api)
        elif option == "think":
            self.think = not self.think
            await api.set_deepthink(self.think)
            self.refresh_ui()
            console.print(f"[bold #38BDF8]DeepThink reasoning: {'ENABLED' if self.think else 'DISABLED'}[/bold #38BDF8]")
        elif option == "search":
            if self.mode == "expert":
                console.print("[dim #64748B]Smart Search is not available in Expert mode.[/dim #64748B]")
            else:
                self.search = not self.search
                self.refresh_ui()
                console.print(f"[bold #38BDF8]Smart Web Search: {'ENABLED' if self.search else 'DISABLED'}[/bold #38BDF8]")
        elif option == "new":
            await api.new_chat()
            self.is_new_chat = True
            self.refresh_ui()
            console.print("[bold #38BDF8]Session reset successfully.[/bold #38BDF8]")

    def _print_tool_call_panel(self, call: dict):
        tool_name = call.get("tool")
        tool_args = call.get("args", {})
        if tool_name in FILE_VISUAL_TOOL_NAMES or tool_name in ["render_plan", "todo"]:
            return
        if tool_name in ["run_cmd", "run_background_cmd"]:
            cmd_preview = tool_args.get("command", "")
            if len(cmd_preview) > 55:
                cmd_preview = cmd_preview[:52] + "..."
            summary_str = f"{tool_name}: {cmd_preview}"
        elif tool_name in ["read_file", "file_info"]:
            path_preview = tool_args.get("path", "") or "<no path>"
            summary_str = f"{tool_name}: {path_preview}"
            if tool_name == "read_file" and tool_args.get("start_line"):
                summary_str += (
                    f" [{tool_args.get('start_line')}-{tool_args.get('end_line')}]"
                )
        elif tool_name == "run_python":
            code_preview = tool_args.get("code", "").replace("\n", " ")
            if len(code_preview) > 55:
                code_preview = code_preview[:52] + "..."
            summary_str = f"run_python: {code_preview}"
        else:
            arg_preview = str(tool_args)
            if len(arg_preview) > 45:
                arg_preview = arg_preview[:42] + "..."
            summary_str = f"{tool_name}: {arg_preview}"

        console.print(Panel(
            f"[bold #536DFE]> {summary_str}[/bold #536DFE]",
            box=box.ROUNDED,
            border_style="#536DFE",
            padding=(0, 1),
            expand=False,
        ))

    def _file_visual_snapshot(self, path: str):
        abs_path = os.path.abspath(os.path.join(self.tools.cwd, path or ""))
        exists = os.path.isfile(abs_path)
        if not exists:
            return False, []
        if Path(abs_path).suffix.lower() in {
            ".xlsx", ".xls", ".docx", ".pptx", ".zip"
        }:
            return True, []
        try:
            with open(abs_path, "r", encoding="utf-8", errors="replace") as file:
                return True, file.read().splitlines()
        except OSError:
            return True, []

    def _file_visual_style(self, action: str):
        return {
            "created": ("FILE CREATED", "#10B981"),
            "modified": ("FILE MODIFIED", "#38BDF8"),
            "deleted": ("FILE DELETED", "#EF4444"),
        }[action]

    def _file_visual_line(
        self,
        action: str,
        filename: str,
        spinner: str = "",
        added: int = 0,
        removed: int = 0,
    ):
        label, color = self._file_visual_style(action)
        line = Text()
        line.append("▌", style=f"bold {color}")
        line.append(" ")
        line.append(label.ljust(13), style=f"bold {color}")
        line.append(" ")
        line.append(filename or "<unknown>", style="bold white")
        if spinner:
            line.append(f" {spinner}", style=f"bold {color}")
        else:
            if added:
                line.append(f" +{added}", style="bold #10B981")
            if removed:
                line.append(f" -{removed}", style="bold #EF4444")
        return line

    def _file_visual_diff(self, before_lines: list, after_lines: list):
        added = 0
        removed = 0
        matcher = difflib.SequenceMatcher(a=before_lines, b=after_lines, autojunk=False)
        for operation, before_start, before_end, after_start, after_end in matcher.get_opcodes():
            if operation in ("insert", "replace"):
                added += after_end - after_start
            if operation in ("delete", "replace"):
                removed += before_end - before_start
        return added, removed

    def _detect_streaming_file_activities(self, text: str):
        """Find file tool calls as soon as their tool name and path are streamed."""
        tool_pattern = re.compile(
            r'"tool"\s*:\s*"('
            + "|".join(sorted(FILE_VISUAL_TOOL_NAMES))
            + r')"'
        )
        matches = list(tool_pattern.finditer(text or ""))
        detected = []
        for index, match in enumerate(matches):
            segment_end = (
                matches[index + 1].start()
                if index + 1 < len(matches)
                else min(len(text), match.end() + 4000)
            )
            segment = text[match.end():segment_end]
            path_match = re.search(
                r'"path"\s*:\s*"((?:\\.|[^"\\])*)"',
                segment,
                flags=re.DOTALL,
            )
            if not path_match:
                continue
            raw_path = path_match.group(1)
            try:
                path = json.loads(f'"{raw_path}"')
            except Exception:
                path = raw_path.replace(r"\\", "\\").replace(r"\"", '"')
            path = str(path or "").strip()
            if not path:
                continue
            tool_name = match.group(1)
            before_exists, _ = self._file_visual_snapshot(path)
            action = (
                "created"
                if tool_name != "edit_file" and not before_exists
                else "modified"
            )
            detected.append((tool_name, path, action))
        return detected

    def _streaming_file_renderable(self, frame_index: int):
        renderables = []
        for activity in self._stream_file_activities.values():
            renderables.append(self._file_visual_line(
                activity["action"],
                activity["filename"],
                spinner=FILE_PROGRESS_FRAMES[
                    frame_index % len(FILE_PROGRESS_FRAMES)
                ],
            ))
        return Group(*renderables)

    def _update_streaming_file_visual(
        self,
        text: str,
        frame_index: int,
        leading_newlines: int = 0,
    ):
        detected = self._detect_streaming_file_activities(text)
        if not detected:
            return set()

        if not hasattr(self, "_stream_file_activities"):
            self._stream_file_activities = {}
        for tool_name, path, action in detected:
            key = (tool_name, path)
            self._stream_file_activities.setdefault(key, {
                "action": action,
                "filename": os.path.basename(path) or path,
            })

        renderable = self._streaming_file_renderable(frame_index)
        if getattr(self, "_active_file_live", None) is None:
            if leading_newlines > 0:
                sys.stdout.write("\n" * leading_newlines)
                sys.stdout.flush()
            self._active_file_live = Live(
                renderable,
                console=console,
                refresh_per_second=20,
                transient=True,
            )
            self._active_file_live.start()
        else:
            self._active_file_live.update(renderable, refresh=True)
        return {(tool_name, path) for tool_name, path, _ in detected}

    def _stop_streaming_file_visual(self):
        live = getattr(self, "_active_file_live", None)
        if live is not None:
            try:
                live.stop()
            finally:
                self._active_file_live = None
        if hasattr(self, "_stream_file_activities"):
            self._stream_file_activities = {}

    async def _execute_file_tool_call(self, tool_name: str, tool_args: dict):
        tool_args = dict(tool_args or {})
        stream_previewed = bool(tool_args.pop("_stream_previewed", False))
        path = str(tool_args.get("path", "") or "")
        filename = os.path.basename(path) or path or "<unknown>"
        before_exists, before_lines = self._file_visual_snapshot(path)
        pending_action = (
            "created"
            if tool_name in ("write_file", "make_excel") and not before_exists
            else "modified"
        )
        if stream_previewed:
            result = await self.tools.execute_tool(
                tool_name,
                tool_args,
                allowed_tools=self.allowed_tools,
            )
        else:
            started_at = asyncio.get_running_loop().time()
            frame_index = 0
            operation = asyncio.create_task(self.tools.execute_tool(
                tool_name,
                tool_args,
                allowed_tools=self.allowed_tools,
            ))

            with Live(
                self._file_visual_line(
                    pending_action,
                    filename,
                    spinner=FILE_PROGRESS_FRAMES[0],
                ),
                console=console,
                refresh_per_second=20,
                transient=True,
            ) as live:
                while (
                    not operation.done()
                    or asyncio.get_running_loop().time() - started_at
                    < FILE_PROGRESS_MIN_SECONDS
                ):
                    live.update(self._file_visual_line(
                        pending_action,
                        filename,
                        spinner=FILE_PROGRESS_FRAMES[
                            frame_index % len(FILE_PROGRESS_FRAMES)
                        ],
                    ))
                    frame_index += 1
                    await asyncio.sleep(FILE_PROGRESS_FRAME_SECONDS)
                result = await operation

        result_text = str(result)
        if result_text.startswith(("[Error", "[Access denied")):
            return result

        after_exists, after_lines = self._file_visual_snapshot(path)
        if before_exists and not after_exists:
            final_action = "deleted"
        elif not before_exists and after_exists:
            final_action = "created"
        else:
            final_action = "modified"
        added, removed = self._file_visual_diff(before_lines, after_lines)
        console.print(self._file_visual_line(
            final_action,
            filename,
            added=added,
            removed=removed,
        ))
        return result

    async def _execute_one_tool_call(self, call: dict):
        tool_name = call.get("tool")
        tool_args = call.get("args", {})
        try:
            if tool_name in FILE_VISUAL_TOOL_NAMES:
                result = await self._execute_file_tool_call(tool_name, tool_args)
            else:
                result = await self.tools.execute_tool(
                    tool_name,
                    tool_args,
                    allowed_tools=self.allowed_tools,
                )
        except Exception as error:
            result = f"[Error executing tool '{tool_name}': {type(error).__name__}: {error}]"
        clean_result = sanitize_tool_result(result)
        return f"[TOOL RESULT for '{tool_name}']:\n{clean_result}"

    async def _execute_tool_calls(self, tool_calls: list):
        outputs = []
        executed_count = 0
        index = 0

        while index < len(tool_calls):
            call = tool_calls[index]
            tool_name = call.get("tool")
            if tool_name in PARALLEL_READ_ONLY_TOOLS:
                end = index + 1
                while (
                    end < len(tool_calls)
                    and tool_calls[end].get("tool") in PARALLEL_READ_ONLY_TOOLS
                ):
                    end += 1
                batch = tool_calls[index:end]
                for batch_call in batch:
                    self._print_tool_call_panel(batch_call)
                self.tools.flush_todos()
                batch_outputs = await asyncio.gather(
                    *(self._execute_one_tool_call(batch_call) for batch_call in batch)
                )
                outputs.extend(batch_outputs)
                executed_count += len(batch)
                index = end
                continue

            self._print_tool_call_panel(call)
            if tool_name != "todo":
                self.tools.flush_todos()
            outputs.append(await self._execute_one_tool_call(call))
            if tool_name != "todo":
                executed_count += 1
            index += 1

        self.tools.flush_todos()
        return outputs, executed_count

    async def process_query(self, user_input: str, api: DeepAPI):
                                                                                                                    
        
        original_goal = user_input
        failure_counts = {}
        pending_verification = False
        strategy_change_required = False
        verification_nudges = 0
        current_prompt = user_input
                                                                                      
        if self.is_new_chat:
            system_prompt = build_agent_system_prompt(self.agent_style, self.think, self.tools.cwd)
            if self.mode == "expert":
                system_prompt += (
                    "\n\n[EXPERT MODE]: Never invoke or refer to DeepSeek's internal "
                    "Search/link-reading feature and never tell the user to switch to "
                    "Fast/Instant mode. Use only the external tools listed in this prompt."
                )
            current_prompt = (
                f"{system_prompt}\n\n[ПОЛЬЗОВАТЕЛЬСКИЙ ЗАПРОС]:\n"
                f"{user_input}"
            )
            self.is_new_chat = False

        if self.is_new_chat:
            await api.new_chat()
            self.is_new_chat = False

        await api.set_mode(self.mode)
        if self.mode == "instant":
            await api.set_search(self.search)
        await api.set_deepthink(self.think)

        loop_count = 0
        did_execute_tool = False
        executed_tool_count = 0
        todo_nudged = False
        empty_responses = 0
                                                                                            
                                                              
        if self.tools.todos and all(
            item["status"] in ("done", "cancelled") for item in self.tools.todos
        ):
            self.tools.todos = []

        spinner_frames = ["|", "/", "-", "\\"]

        while True:
            loop_count += 1

                                                                                              
            if loop_count == 1:
                sys.stdout.write("\n")
                sys.stdout.flush()

            try:
                stream_gen = await api.send_message(current_prompt, stream=True)
            except Exception as send_error:
                empty_responses += 1
                message = str(send_error).replace("\n", " ")[:200]
                                                                                        
                console.print(
                    f"\nЗапрос не ушёл ({type(send_error).__name__}): {message}",
                    style="bold #F59E0B",
                    markup=False,
                    highlight=False,
                )
                if empty_responses > MAX_EMPTY_RESPONSE_RETRIES:
                    console.print("[bold red]Повторы исчерпаны. Останавливаюсь.[/bold red]")
                    break
                await asyncio.sleep(2 * empty_responses)
                continue

            printed_answer = ""
            thinking_line_printed = False
            answer_started = False
            has_printed_text_this_turn = False
            final_think = ""
            final_answer = ""
            stream_previewed_file_calls = set()
            stream_visual_frame = 0
            stream_file_visual_was_shown = False
            self._stop_streaming_file_visual()

            async for chunk in stream_gen:
                if not isinstance(chunk, dict):
                    continue

                final_think = chunk.get("think", "")
                final_answer = chunk.get("answer", "")

                if final_answer:
                    answer_started = True

                if final_think and not answer_started:
                    thinking_line_printed = True
                    tok_count = count_tokens(final_think)
                    frame = spinner_frames[len(final_think) % len(spinner_frames)]
                    sys.stdout.write(f"\r\033[K\033[38;2;83;109;254m[THINKING {frame}] Reasoning in progress... ({tok_count} tokens)\033[0m")
                    sys.stdout.flush()
                
                if final_answer:
                    if thinking_line_printed:
                        sys.stdout.write("\r\033[K")
                        sys.stdout.flush()
                        thinking_line_printed = False

                    cleaned = render_terminal_markup(get_clean_text(final_answer), final=False)
                    
                    if len(cleaned) > len(printed_answer):
                        if cleaned.startswith(printed_answer):
                            new_chars = cleaned[len(printed_answer):]
                        else:
                                                                        
                            common_len = 0
                            min_l = min(len(cleaned), len(printed_answer))
                            for i in range(min_l):
                                if cleaned[i] == printed_answer[i]:
                                    common_len += 1
                                else:
                                    break
                            new_chars = cleaned[common_len:]
                        
                                                                                     
                        if not has_printed_text_this_turn:
                            if loop_count > 1 and did_execute_tool:
                                sys.stdout.write("\n")
                            has_printed_text_this_turn = True
                        await write_streaming_chars(new_chars)
                        printed_answer = cleaned

                    detected_file_calls = self._update_streaming_file_visual(
                        final_answer,
                        stream_visual_frame,
                        leading_newlines=(
                            max(
                                0,
                                2 - (
                                    len(printed_answer)
                                    - len(printed_answer.rstrip("\n"))
                                ),
                            )
                            if printed_answer
                            and not stream_file_visual_was_shown
                            else 0
                        ),
                    )
                    if detected_file_calls:
                        stream_previewed_file_calls.update(detected_file_calls)
                        stream_visual_frame += 1
                        stream_file_visual_was_shown = True

                                                                                       
            self._stop_streaming_file_visual()
            final_display = render_terminal_markup(get_clean_text(final_answer), final=True)
            if len(final_display) > len(printed_answer):
                if final_display.startswith(printed_answer):
                    await write_streaming_chars(final_display[len(printed_answer):])
                else:
                    common_len = 0
                    for old_char, new_char in zip(printed_answer, final_display):
                        if old_char != new_char:
                            break
                        common_len += 1
                    await write_streaming_chars(final_display[common_len:])
                printed_answer = final_display

            if printed_answer and not stream_file_visual_was_shown:
                nl_count = len(printed_answer) - len(printed_answer.rstrip('\n'))
                if nl_count == 0:
                    console.print("\n")
                elif nl_count == 1:
                    console.print()

            final_answer = final_answer.strip()
            final_think = final_think.strip()

                                                                                       
                                                                                 
                                                                                        
                                                                      
            if not final_answer:
                if final_think:
                    console.print(
                        "\n[bold #F59E0B]Ответ оборван: модель начала размышлять, но текста "
                        "ответа нет.[/bold #F59E0B]"
                    )
                    console.print(
                        "[dim #94A3B8]Автопродолжение не выполняется. Напиши, что делать "
                        "дальше.[/dim #94A3B8]"
                    )
                    break

                empty_responses += 1
                if empty_responses > MAX_EMPTY_RESPONSE_RETRIES:
                    console.print(
                        f"\n[bold red]Пустой ответ {empty_responses} раз(а) подряд. "
                        "Останавливаюсь.[/bold red]"
                    )
                    break

                delay = 2 * empty_responses
                console.print(
                    f"\n[bold #F59E0B]Пустой ответ от API (попытка {empty_responses} из "
                    f"{MAX_EMPTY_RESPONSE_RETRIES}). Повтор через {delay} с...[/bold #F59E0B]"
                )
                await asyncio.sleep(delay)
                continue

            empty_responses = 0

            full_rendered = ""
            if final_think:
                full_rendered += f"--- [РАЗМЫШЛЕНИЯ] ---\n{get_clean_text(final_think)}\n\n"
            if final_answer:
                full_rendered += get_clean_text(final_answer)

            self.last_response = full_rendered

                                                             
            tool_calls = extract_tool_calls(final_answer)
            if tool_calls:
                    for call in tool_calls:
                        tool_name = str(call.get("tool", ""))
                        tool_path = str(call.get("args", {}).get("path", "") or "")
                        if (tool_name, tool_path) in stream_previewed_file_calls:
                            call.setdefault("args", {})["_stream_previewed"] = True
                    did_execute_tool = True
                    tool_outputs, current_executed_count = await self._execute_tool_calls(
                        tool_calls
                    )
                    executed_tool_count += current_executed_count

                    last_mutation_index = -1
                    for call_index, (call, output) in enumerate(
                        zip(tool_calls, tool_outputs)
                    ):
                        tool_name = str(call.get("tool", "unknown"))
                        if _tool_result_failed(output):
                            failure_key = (
                                tool_name,
                                json.dumps(
                                    call.get("args", {}),
                                    ensure_ascii=False,
                                    sort_keys=True,
                                    default=str,
                                ),
                            )
                            failure_counts[failure_key] = (
                                failure_counts.get(failure_key, 0) + 1
                            )
                            if failure_counts[failure_key] >= 2:
                                strategy_change_required = True
                        else:
                            strategy_change_required = False

                        if tool_name in MUTATING_TOOLS_REQUIRING_VERIFICATION:
                            last_mutation_index = call_index

                    if last_mutation_index >= 0:
                        pending_verification = not any(
                            str(call.get("tool", "")) in VERIFICATION_TOOLS
                            and call_index > last_mutation_index
                            and not _tool_result_failed(tool_outputs[call_index])
                            for call_index, call in enumerate(tool_calls)
                        )
                    elif pending_verification:
                        pending_verification = not any(
                            str(call.get("tool", "")) in VERIFICATION_TOOLS
                            and not _tool_result_failed(tool_outputs[call_index])
                            for call_index, call in enumerate(tool_calls)
                        )

                                                                                
                    feedback = "\n\n".join(tool_outputs)
                    if len(feedback) > MAX_TOOL_FEEDBACK_CHARS:
                        feedback = (
                            feedback[:MAX_TOOL_FEEDBACK_CHARS]
                            + f"\n\n[Truncated: tool output exceeded {MAX_TOOL_FEEDBACK_CHARS} chars. "
                            "Request the remaining part in a follow-up call.]"
                        )
                                                                                                
                                                                                               
                                                                                                 
                    if "todo" in self.allowed_tools and not todo_nudged:
                        if executed_tool_count >= TODO_NUDGE_AFTER_TOOLS and not self.tools.todos:
                            todo_nudged = True
                            feedback += (
                                "\n\n[СИСТЕМА]: задача оказалась многошаговой, а список задач не заведён. "
                                "Следующим действием вызови todo с action=set и распиши оставшиеся шаги, "
                                "затем продолжай работу, отмечая их по мере выполнения."
                            )
                        elif self.tools.todos and not any(
                            item["status"] == "active" for item in self.tools.todos
                        ) and any(item["status"] == "pending" for item in self.tools.todos):
                            todo_nudged = True
                            feedback += (
                                "\n\n[СИСТЕМА]: в списке задач нет ни одной активной. Отметь "
                                "выполненные через complete и переведи текущую в active."
                            )

                    task_state = build_agent_task_state(
                        original_goal,
                        self.tools.todos,
                        pending_verification=pending_verification,
                        strategy_change_required=strategy_change_required,
                        include_goal=(
                            loop_count % TASK_GOAL_REMINDER_INTERVAL == 0
                        ),
                    )
                    current_prompt = (
                        task_state
                        + ("\n" if task_state else "")
                        + "[РЕЗУЛЬТАТЫ ИНСТРУМЕНТОВ]\n"
                        + feedback
                        + "\n[КОНЕЦ РЕЗУЛЬТАТОВ]\n\n"
                        "Продолжай исходную задачу. Если она подтверждённо выполнена — "
                        "дай финальный ответ; иначе выбери следующий проверяемый шаг."
                    )
                    continue
            else:
                if pending_verification and verification_nudges < 2:
                    verification_nudges += 1
                    current_prompt = (
                        build_agent_task_state(
                            original_goal,
                            self.tools.todos,
                            pending_verification=True,
                            strategy_change_required=strategy_change_required,
                            include_goal=(
                                loop_count % TASK_GOAL_REMINDER_INTERVAL == 0
                            ),
                        )
                        + "\nТы попытался завершить задачу до проверки изменённого "
                        "артефакта. Следующим сообщением вызови релевантный инструмент "
                        "проверки. Не повторяй неподтверждённый финальный ответ."
                    )
                    continue
                break

                                                                                
             
                                                                                

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="DEEPX AGENT: Autonomous AI Engineering CLI")
    parser.add_argument("--headful", action="store_true", help="Launch browser in visible mode")
    parser.add_argument("--state", type=str, default="state.json", help="Path to playwright state file")
    args = parser.parse_args()

    app = DeepCLIApp(state_file=args.state, headless=not args.headful)
    asyncio.run(app.run())
